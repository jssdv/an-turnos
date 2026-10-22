from datetime import date, timedelta

import capacitacion
import db
import exportar
import flujo
import tempfile
import openpyxl
import puestos
import recargos
import reglas
import solver
from dominio import Colaborador, Novedad, Turno, festivos


P = flujo.periodo_mes(2026, 10)
K = flujo.clave(*P)  # «2026-10-01/2026-10-31»

def cuadro_vacio(bd):
    return bd.contexto(2026)


def test_festivos():
    f = festivos(2026)
    assert date(2026, 10, 12) in f and date(2026, 11, 2) in f and date(2026, 8, 17) in f
    assert date(2026, 4, 2) in f and date(2026, 4, 3) in f  # Semana Santa (Pascua 5-abr)
    assert date(2026, 5, 18) in f and date(2026, 6, 8) in f and date(2026, 6, 15) in f
    assert len(f) == 18


def test_turnos_y_pausas():
    t = Turno("C", "Corrido", 360, 1080)
    assert t.duracion == 720 and t.horas == 12
    t.pausas = ({"ini": 360, "min": 60, "pagada": False},)
    assert t.horas == 11
    t.pausas = ({"ini": 360, "min": 60, "pagada": True},)
    assert t.horas == 12
    assert Turno("NC", "", 1080, 360).duracion == 720


def test_recargos():
    bd = db.BD()
    ctx, tar = bd.contexto(2026), bd.tarifas()
    # NC dominical 18:00-06:00 (domingo 18-oct-2026): 1 h antes de las 19 ya es noche → 18:00 diurno hasta 19:00
    segs = list(recargos.segmentos(ctx.turnos["NC"], date(2026, 10, 18), ctx, tar))
    horas = lambda cond: sum((b - a).total_seconds() / 3600 for a, b, n, f in segs if cond(n, f))
    assert horas(lambda n, f: n) == 11 and horas(lambda n, f: not n) == 1
    assert horas(lambda n, f: f) == 6  # 18:00-24:00 domingo; el lunes de 0 a 6 es hábil
    # Tarifas por vigencia
    assert recargos.tarifa(tar, "dominical", date(2026, 6, 30)) == 80
    assert recargos.tarifa(tar, "dominical", date(2026, 7, 1)) == 90
    assert recargos.tarifa(tar, "dominical", date(2027, 7, 1)) == 100
    assert recargos.tarifa(tar, "noche_inicio", date(2025, 12, 24)) == 1260
    assert recargos.tarifa(tar, "noche_inicio", date(2025, 12, 25)) == 1140
    # Pausa no pagada de 60 min en C baja 12→11 h y esa hora no recibe recargo
    c = ctx.turnos["C"]
    c.pausas = ({"ini": 360, "min": 60, "pagada": False},)
    assert sum((b - a).total_seconds() / 3600 for a, b, n, f in recargos.segmentos(c, date(2026, 10, 20), ctx, tar)) == 11


def test_reporte_y_ventana():
    bd = db.BD()
    p = Colaborador(0, "A")
    p.id = bd.guardar_colaborador(p)
    ctx, tar = bd.contexto(2026), bd.tarifas()
    p = ctx.colabs[0]
    fechas = [date(2026, 10, 1) + timedelta(days=i) for i in range(21)]  # 3 semanas
    cuadro = {(p.id, f): "M" for f in fechas[:7]}  # 7 × 6 = 42 h
    r = recargos.reporte(ctx, tar, p, cuadro, fechas, 3)
    assert r["trabajadas"] == 42 and r["contratadas"] == 126 and r["pendientes"] == 84 and r["adicionales"] == 0
    cuadro = {(p.id, f): "C" for f in fechas[:11]}  # 132 h en 126 → 6 h extra
    r = recargos.reporte(ctx, tar, p, cuadro, fechas, 3)
    assert r["adicionales"] == 6 and r["cat"][("extra", "diurna", "hábil")] + r["cat"][("extra", "diurna", "festivo")] == 6
    # Novedad: 7 días de vacaciones bajan lo contratado en 42 h
    ctx2 = bd.contexto(2026)
    ctx2.novedades.append(Novedad(p.id, "Vacaciones", fechas[0], fechas[6]))
    ctx2.__post_init__()
    assert recargos.reporte(ctx2, tar, p, {}, fechas, 3)["contratadas"] == 84


def test_reglas():
    bd = db.BD()
    ids = [bd.guardar_colaborador(Colaborador(0, f"P{i}")) for i in range(2)]
    ctx = bd.contexto(2026)
    a, b = ctx.colabs
    lunes = date(2026, 10, 19)
    cuadro = {(a.id, lunes): "NC", (a.id, lunes + timedelta(days=1)): "M"}  # NC termina 6:00, M empieza 6:00
    al = reglas.validar(ctx, cuadro, [lunes, lunes + timedelta(days=1)])
    assert any("menos de 12 h" in x.texto for x in al)
    assert any("faltan" in x.texto for x in al)  # sin cobertura
    # Persona restringida y estudio
    r = Colaborador(0, "R", tipo="restringido", permitidos=("M", "T"), estudio=((0, "M"),))
    r.id = bd.guardar_colaborador(r)
    ctx = bd.contexto(2026)
    r = ctx.colabs[-1]
    assert ctx.bloqueo(r, ctx.turnos["NC"], lunes) == "turno no permitido"
    assert ctx.bloqueo(r, ctx.turnos["M"], lunes) == "estudio"
    assert ctx.bloqueo(r, ctx.turnos["T"], lunes) is None
    assert ctx.bloqueo(a, ctx.turnos["C"], date(2026, 10, 18)) == "turno no permitido en domingo/festivo"
    assert reglas.validar_catalogo(ctx.turnos)  # C y NC sin pausa → aviso


def test_libre_tras_noche():
    bd = db.BD()
    for i in range(2):
        bd.guardar_colaborador(Colaborador(0, f"P{i}"))
    ctx = bd.contexto(2026)
    assert ctx.libre_tras_noche == 1
    a = ctx.colabs[0]
    lunes, martes = date(2026, 10, 19), date(2026, 10, 20)
    cuadro = {(a.id, lunes): "NC", (a.id, martes): "T"}  # T empieza 12:00: cumple las 12 h, pero no hay día libre
    sin = lambda **kw: [x.texto for x in reglas.validar(ctx, cuadro, [lunes, martes], **kw) if "sin día libre" in x.texto]
    assert len(sin()) == 1
    assert sin(fijos={(a.id, martes)}) == []  # decisión manual: no se reclama
    bd.set_param("libre_tras_noche", 0)  # regla apagada: no hay alerta
    ctx = bd.contexto(2026)
    assert ctx.libre_tras_noche == 0 and sin() == []


def test_fin_de_semana_y_libres():
    bd = db.BD()
    bd.guardar_colaborador(Colaborador(0, "A"))
    ctx = bd.contexto(2026)
    a = ctx.colabs[0]
    assert (ctx.libres_semana, ctx.fines_libres) == (1, 1)
    sab, dom = date(2026, 10, 3), date(2026, 10, 4)
    mes = [date(2026, 10, 1) + timedelta(days=i) for i in range(10)]  # 1-oct (jue) a 10-oct (sáb): un solo fin de semana completo
    sin_fin = lambda cuadro, **kw: [x.texto for x in reglas.validar(ctx, cuadro, mes, **kw) if "fin de semana" in x.texto]
    assert len(sin_fin({(a.id, sab): "M"})) == 1
    assert sin_fin({(a.id, sab): "M"}, fijos={(a.id, sab)}) == []  # decisión manual
    assert sin_fin({(a.id, d): "M" for d in mes if d not in (sab, dom)}) == []  # sábado y domingo libres
    semana = [date(2026, 10, 5) + timedelta(days=i) for i in range(7)]
    todo = {(a.id, d): "M" for d in semana}
    assert any("día(s) libre(s)" in x.texto for x in reglas.validar(ctx, todo, semana))


def test_etiqueta_de_puesto():
    import puestos
    assert puestos.etiqueta("M", 2) == "M·2" and puestos.etiqueta("T", 0) == "T·A" and puestos.etiqueta("NC", None) == "NC"
    filas = [(date(2026, 10, 5), "M", 3, 7, 1), (date(2026, 10, 5), "T", 3, 7, 1), (date(2026, 10, 5), "T", 1, 8, 2)]
    assert puestos.por_persona(filas) == {(7, date(2026, 10, 5)): 3, (8, date(2026, 10, 5)): 1}


def test_rachas():
    bd = db.BD()
    bd.guardar_colaborador(Colaborador(0, "A"))
    ctx = bd.contexto(2026)
    a = ctx.colabs[0]
    dias = [date(2026, 10, 5) + timedelta(days=i) for i in range(7)]
    avisos = lambda cuadro, **kw: [x.texto for x in reglas.validar(ctx, cuadro, dias, **kw) if "seguidos" in x.texto or "seguidas" in x.texto]
    tres = {(a.id, dias[i]): "NC" for i in (0, 2, 4)}  # NC L NC L NC: 3 noches en racha, máximo 2
    assert len(avisos(tres)) == 1 and "3 noches" in avisos(tres)[0]
    assert avisos({(a.id, dias[i]): "NC" for i in (0, 2)}) == []
    assert avisos(tres, fijos={(a.id, dias[4])}) == []  # edición manual
    cuatro = {(a.id, dias[i]): "M" for i in range(4)}  # 4 M seguidas, máximo 3
    assert len(avisos(cuatro)) == 1 and "4 M" in avisos(cuatro)[0]
    assert avisos({(a.id, dias[i]): "M" for i in range(3)}) == [] and avisos(cuatro, fijos={(a.id, dias[3])}) == []
    assert len(avisos({(a.id, dias[i]): "C" for i in (0, 1, 2)})) == 1 and avisos({(a.id, dias[i]): "C" for i in (0, 1)}) == []


def test_periodos():
    ini, fin = flujo.periodo_mes(2026, 2)
    assert (ini, fin) == (date(2026, 2, 1), date(2026, 2, 28)) and len(flujo.fechas(ini, fin)) == 28
    lunes = date(2026, 10, 5)
    assert lunes.weekday() == 0 and flujo.periodo_semanas(lunes, 4) == (lunes, date(2026, 11, 1))
    assert flujo.desde_clave(flujo.clave(lunes, date(2026, 11, 1))) == (lunes, date(2026, 11, 1))
    assert flujo.etiqueta(*P) == "Octubre 2026" and flujo.etiqueta(lunes, date(2026, 11, 1)) == "05-10 AL 01-11 2026"


def test_migracion_v1():
    """Una base con claves «AAAA-MM» pasa a claves de periodo sin perder el cuadro."""
    import sqlite3
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        ruta = tmp + "/v1.db"
        bd = db.BD(ruta)
        bd.c.execute("INSERT INTO asignacion VALUES ('2026-02','cambios',1,'2026-02-03','M','solver',0)")
        bd.c.execute("INSERT INTO puesto_dia VALUES ('2026-02','2026-02-03','M',2,1,1)")
        bd.c.execute("DELETE FROM parametro WHERE clave IN ('periodo_modo','periodo_semanas')")
        bd.c.execute("PRAGMA user_version = 1")
        bd.c.commit()
        bd.c.close()
        bd = db.BD(ruta)
        k = "2026-02-01/2026-02-28"
        assert bd.periodos() == [k] and bd.asignaciones(k) == {(1, date(2026, 2, 3)): "M"} and len(bd.puestos(k)) == 1
        assert bd.param("periodo_modo") == "mes" and bd.c.execute("PRAGMA user_version").fetchone()[0] == 4


def test_esquema_por_semanas():
    """4 semanas desde un lunes: el cuadro respeta los descansos con el periodo anterior y no usa el mes."""
    bd = db.BD()
    for i in range(13):
        bd.guardar_colaborador(Colaborador(0, f"P{i}"))
    lunes = date(2026, 10, 5)
    ini, fin = flujo.periodo_semanas(lunes, 4)
    # el periodo anterior terminó el domingo con una noche de la persona 1: el lunes 5 debe quedar libre
    bd.guardar_cuadro(flujo.clave(lunes - timedelta(days=28), ini - timedelta(days=1)), {(1, ini - timedelta(days=1)): "NC"})
    info = flujo.generar(bd, ini, fin, tiempo=10)
    k = flujo.clave(ini, fin)
    cuadro, dias = bd.asignaciones(k), flujo.fechas(ini, fin)
    assert len(dias) == 28 and dias[0].weekday() == 0 and dias[-1].weekday() == 6 and not info["faltan"]
    assert (1, ini) not in cuadro  # libre tras la noche del periodo anterior
    ctx = bd.contexto(2026)
    assert not [a for a in reglas.validar(ctx, cuadro, dias) if a.nivel == "error"]
    assert bd.puestos(k) and bd.capacitacion(k) and k in bd.periodos()


def test_banco_de_horas():
    bd = db.BD()
    a = bd.guardar_colaborador(Colaborador(0, "A"))
    b = bd.guardar_colaborador(Colaborador(0, "B"))
    lunes = date(2026, 10, 5)
    ini, fin = flujo.periodo_semanas(lunes, 1)
    k = flujo.clave(ini, fin)
    dias = flujo.fechas(ini, fin)
    # A trabaja 4 corridos (48 h) y B 6 mañanas (36 h) frente a 42 h contratadas
    bd.guardar_cuadro(k, {**{(a, d): "C" for d in dias[:4]}, **{(b, d): "M" for d in dias[:6]}})
    saldo = lambda filas, nombre, clave: next(f[clave] for f in filas if f["nombre"] == nombre)
    abierto = flujo.estado_banco(bd, ini, fin)  # proyectado mientras está abierto
    assert saldo(abierto, "A", "periodo") == 6 and saldo(abierto, "B", "periodo") == -6 and not abierto[0]["cerrado"]
    flujo.cerrar(bd, ini, fin)
    assert bd.periodo_cerrado(k) and bd.cierres() == [(k, a, 42.0, 48.0), (k, b, 42.0, 36.0)]
    cerrado = flujo.estado_banco(bd, ini, fin)
    assert saldo(cerrado, "A", "total") == 6 and saldo(cerrado, "A", "anterior") == 0 and cerrado[0]["cerrado"]
    # la semana siguiente arranca con el saldo anterior
    ini2, fin2 = flujo.periodo_semanas(lunes + timedelta(days=7), 1)
    sig = flujo.estado_banco(bd, ini2, fin2)
    assert saldo(sig, "A", "anterior") == 6 and saldo(sig, "B", "anterior") == -6 and saldo(sig, "A", "periodo") == 0
    # movimientos manuales: compensar 4 h a A y registrar un saldo inicial de +10 a B
    mov = bd.agregar_movimiento(a, date(2026, 10, 14), -4, "Compensado en descanso", "viernes libre")
    bd.agregar_movimiento(b, date(2026, 10, 1), 10, "Saldo inicial")
    sig = flujo.estado_banco(bd, ini2, fin2)
    assert saldo(sig, "A", "total") == 2 and saldo(sig, "B", "total") == 4
    assert [m[3] for m in bd.movimientos(a)] == [-4] and flujo.CONCEPTOS_BANCO["Pagado en nómina"] is True
    bd.borrar_movimiento(mov)
    assert saldo(flujo.estado_banco(bd, ini2, fin2), "A", "total") == 6
    # historial: un periodo cerrado, 84 h trabajadas y saldo total 0 (+6 y −6)
    h = flujo.historial(bd)
    assert len(h) == 1 and h[0]["cerrado"] and h[0]["trabajadas"] == 84 and h[0]["saldo"] == 0 and h[0]["personas"] == 2
    bd.reabrir_periodo(k)
    assert not bd.periodo_cerrado(k) and bd.cierres() == [] and flujo.historial(bd)[0]["saldo"] is None
    try:
        flujo.cerrar(bd, ini2, fin2)  # sin cuadro
        assert False
    except ValueError:
        pass


def test_compensar_banco():
    bd = db.BD()
    a = bd.guardar_colaborador(Colaborador(0, "A"))
    b = bd.guardar_colaborador(Colaborador(0, "B"))
    c = bd.guardar_colaborador(Colaborador(0, "C"))
    lunes = date(2026, 10, 5)
    ini, fin = flujo.periodo_semanas(lunes, 1)
    dias = flujo.fechas(ini, fin)
    # A trabajó +6, B −6 y C quedó al día; luego se cierra
    bd.guardar_cuadro(flujo.clave(ini, fin), {**{(a, d): "C" for d in dias[:4]}, **{(b, d): "M" for d in dias[:6]}, **{(c, d): "T" for d in dias[:7]}})
    flujo.cerrar(bd, ini, fin)
    sig = flujo.periodo_semanas(lunes + timedelta(days=7), 1)
    assert flujo.ajustes_banco(bd, *sig) == {a: -6, b: 6}  # a quien se le deben se le baja; a quien debe se le sube
    bd.set_param("banco_tope_favor_h", 4); bd.set_param("banco_tope_deuda_h", 0)
    assert flujo.ajustes_banco(bd, *sig) == {a: -4}  # tope por periodo y dirección apagada
    bd.agregar_movimiento(a, date(2026, 10, 14), -6, "Compensado en descanso")  # ya se compensó: saldo 0
    assert flujo.ajustes_banco(bd, *sig) == {}
    # el periodo anterior al cierre no se ajusta con su propio saldo
    assert flujo.ajustes_banco(bd, ini, fin) == {}
    # el solver baja la meta de A y sube la de B (13 personas, un mes): A termina con menos horas que B
    bd2 = db.BD()
    for i in range(13):
        bd2.guardar_colaborador(Colaborador(0, f"P{i}"))
    ctx = bd2.contexto(2026)
    mes = flujo.fechas(*P)
    cuadro, info = solver.generar(ctx, mes, tiempo=15, ajuste={1: -24, 2: 24})
    horas = lambda pid: sum(ctx.turnos[x].horas for (p, _), x in cuadro.items() if p == pid)
    assert info["ajuste"] == {1: -24, 2: 24} and horas(2) - horas(1) >= 24, (horas(1), horas(2))


def test_cubre_se_calcula_de_los_horarios():
    bd = db.BD()
    assert bd.turnos()["C"].cubre == ("M", "T") and bd.turnos()["NC"].cubre == ("N",)
    bd.guardar_turno(Turno("M7", "Mañana 7", 420, 780, "CCE5FF", ("N",)))  # el valor guardado se ignora
    assert bd.turnos()["M7"].cubre == ("M",)
    t = bd.turnos()["M7"]
    t.inicio, t.fin = 360, 1080
    bd.guardar_turno(t)
    assert bd.turnos()["M7"].cubre == ("M", "T")


def test_versiones():
    bd = db.BD()
    pid = bd.guardar_colaborador(Colaborador(0, "A"))
    f = date(2026, 10, 5)
    bd.guardar_cuadro(K, {(pid, f): "M"})
    bd.set_celda(K, pid, f, "T")
    assert bd.asignaciones(K, "original")[(pid, f)] == "M"
    assert bd.diferencias(K) == [(pid, f, "M", "T")]
    bd.guardar_cuadro(K, {(pid, f): "NC"})  # regenerar respeta la celda fija
    assert bd.asignaciones(K)[(pid, f)] == "T"
    bd.restaurar(K)
    assert bd.diferencias(K) == []
    g = date(2026, 10, 6)  # celda libre en el original: restaurar la deja libre y sin fila
    bd.set_celda(K, pid, g, "M")
    assert bd.diferencias(K) == [(pid, g, "L", "M")]
    bd.restaurar(K, {pid}, {g})
    assert bd.diferencias(K) == [] and (pid, g) not in bd.asignaciones(K)
    bd.set_celda(K, pid, f, "C")
    bd.publicar(K)
    assert bd.diferencias(K) == []
    assert bd.turno_en_uso("C")
    bd.borrar_turno("C")
    assert "C" not in bd.turnos() and "C" in bd.turnos(solo_activos=False)
    # eliminar: sin historial se borra; con historial se desactiva
    nuevo = bd.guardar_colaborador(Colaborador(0, "Sin historia"))
    assert bd.borrar_colaborador(nuevo) == "borrado" and len(bd.colaboradores(solo_activos=False)) == 1
    assert bd.borrar_colaborador(pid) == "desactivado" and bd.colaboradores() == []
    bd.guardar_turno(Turno("X", "X", 0, 360))
    assert bd.borrar_turno("X") == "borrado" and "X" not in bd.turnos(solo_activos=False)


def test_mes_completo():
    """Octubre 2026: 13 personas (1 restringida, 2 estudiantes), novedades; el solver cumple todas las reglas."""
    bd = db.BD()
    for i in range(1, 14):
        p = Colaborador(0, f"P{i}")
        if i == 11:
            p.tipo, p.permitidos, p.puesto_fijo = "restringido", ("M", "T"), 3
        if i == 12:
            p.tipo, p.estudio = "estudiante", ((0, "M"), (3, "M"), (4, "M"), (1, "T"))
        if i == 13:
            p.tipo, p.estudio = "estudiante", ((3, "M"), (4, "M"), (4, "T"))
        bd.guardar_colaborador(p)
    fechas = [date(2026, 10, 1) + timedelta(days=i) for i in range(31)]
    bd.guardar_novedad(Novedad(3, "Vacaciones", date(2026, 10, 12), date(2026, 10, 18)))
    ctx = bd.contexto(2026)
    cuadro, info = solver.generar(ctx, fechas, tiempo=20)
    assert not info["faltan"] and not reglas.validar(ctx, cuadro, fechas)
    assert all((3, date(2026, 10, d)) not in cuadro for d in range(12, 19))
    assert all(cuadro.get((11, f), "M") in ("M", "T") for f in fechas)  # la restringida nunca hace noche
    noche = lambda p, f: cuadro.get((p, f)) == "NC"
    assert not any(noche(p.id, f) and (p.id, f + timedelta(days=1)) in cuadro for p in ctx.colabs for f in fechas)  # tras una noche, día libre
    # racha NC-L-NC-L…: nunca más de `max_noches` (2); corridos: nunca más de `max_c` (2) en días consecutivos
    assert ctx.max_noches == 2 and ctx.max_c == 2
    assert not any(all(noche(p.id, fechas[i + 2 * k]) for k in range(3)) for p in ctx.colabs for i in range(len(fechas) - 4))
    assert not any(all(cuadro.get((p.id, fechas[i + k])) == "C" for k in range(3)) for p in ctx.colabs for i in range(len(fechas) - 2))
    # nunca más de 3 turnos de 6 h iguales seguidos (p. ej. 4 M)
    assert ctx.max_igual == 3
    assert not any(all(cuadro.get((p.id, fechas[i + k])) == cod for k in range(4)) for p in ctx.colabs for cod in ("M", "T") for i in range(len(fechas) - 3))
    # reparto parejo de M, T, C y NC entre quienes pueden hacerlos
    for cod in ("M", "T", "C", "NC"):
        n = [sum(1 for f in fechas if cuadro.get((p.id, f)) == cod) for p in ctx.colabs if p.permitidos is None]
        assert max(n) - min(n) <= 4, (cod, n)
    # una celda fija se respeta al regenerar
    c2, _ = solver.generar(ctx, fechas, fijos={(1, fechas[0]): "NC"}, tiempo=10)
    assert c2[(1, fechas[0])] == "NC" and (1, fechas[1]) not in c2  # la noche fija arrastra su día libre
    # edición manual: forzar un turno el día después de una noche se respeta (y no genera alerta)
    c3, _ = solver.generar(ctx, fechas, fijos={(1, fechas[0]): "NC", (1, fechas[1]): "NC"}, tiempo=10)
    assert c3[(1, fechas[0])] == "NC" and c3[(1, fechas[1])] == "NC"
    # cada persona tiene al menos un fin de semana completo libre (sáb. y dom.)
    fines = [(d, d + timedelta(days=1)) for d in fechas if d.weekday() == 5 and d + timedelta(days=1) in fechas]
    assert all(any((p.id, s) not in cuadro and (p.id, d) not in cuadro for s, d in fines) for p in ctx.colabs)
    # el mes anterior terminó en noche: el día 1 queda libre
    c4, _ = solver.generar(ctx, fechas, previo={(2, date(2026, 9, 30)): "NC"}, tiempo=10)
    assert (2, fechas[0]) not in c4
    # puestos: la restringida siempre en el 3; sin repetir puesto en una misma franja
    filas = puestos.asignar(ctx, cuadro, fechas)
    assert all(p == 3 for f, fr, p, c, g in filas if c == 11)
    assert len({(f, fr, p) for f, fr, p, c, g in filas if p}) == len([1 for f, fr, p, c, g in filas if p])
    # capacitación: lunes, máx. 1 por semana, festivo corre al siguiente hábil (12-oct-2026 es festivo)
    ses, sin = capacitacion.programar(ctx, cuadro, fechas)
    dias = {f for f, _ in ses}
    assert date(2026, 10, 12) not in dias and all(f.weekday() in (0, 1) for f in dias)
    assert len({f - timedelta(days=f.weekday()) for f in dias}) == len(dias)
    assert len(ses) + len(sin) == 13
    # exportación: tres libros, el cuadro marca en amarillo lo cambiado y nunca sobrescribe
    bd.guardar_cuadro(K, cuadro)
    bd.set_celda(K, 1, fechas[0], "NC")
    cam = bd.asignaciones(K)
    with tempfile.TemporaryDirectory() as tmp:
        r1 = exportar.exportar_cuadro(tmp, "Octubre 2026", ctx, cam, bd.asignaciones(K, "original"), fechas,
                                      bd.diferencias(K))
        assert exportar.exportar_cuadro(tmp, "Octubre 2026", ctx, cam, cam, fechas, []).name.endswith("v2.xlsx")
        exportar.exportar_puestos(tmp, "Octubre 2026", ctx, cam, fechas, filas, ses)
        r3 = exportar.exportar_horas(tmp, "Octubre 2026", ctx, bd.tarifas(), cam, fechas, 3)
        wb = openpyxl.load_workbook(r1)
        assert wb.sheetnames == ["Cuadro", "Original", "Cambios"]  # sin puestos calculados no hay hoja de puestos
        r2 = exportar.exportar_cuadro(tmp, "Octubre 2026", ctx, cam, cam, fechas, [], filas)
        wb2 = openpyxl.load_workbook(r2)
        assert wb2.sheetnames[:2] == ["Cuadro", "Cuadro con puestos"]
        ejemplo = next(c.value for fila in wb2["Cuadro con puestos"].iter_rows(min_row=2, min_col=2) for c in fila if c.value and "·" in str(c.value))
        assert ejemplo.split("·")[0] in ("M", "T", "C") and ejemplo.split("·")[1] in "12345A"
        assert wb["Cuadro"].cell(2, 2).fill.fgColor.rgb.endswith("FFFF00") or bd.asignaciones(K, "original").get((1, fechas[0])) == "NC"
        assert openpyxl.load_workbook(r3)["Horas y recargos"].max_row == 14


def test_flujo_y_estudio():
    from dominio import texto_estudio
    assert texto_estudio(((0, "M"), (1, "T"), (3, "M"), (4, "M"), (4, "T"))) == "Lun M, Mar T, Jue M, Vie"
    bd = db.BD()
    for i in range(13):
        bd.guardar_colaborador(Colaborador(0, f"P{i}"))
    info = flujo.generar(bd, *P, tiempo=8)
    assert not info["faltan"] and bd.diferencias(K) == [] and bd.puestos(K) and bd.capacitacion(K)
    # edición manual → aparece como diferencia; «respetar» la conserva y no toca original
    cambios = bd.asignaciones(K)
    f = date(2026, 10, 7)
    bd.set_celda(K, 1, f, "L")
    flujo.generar(bd, *P, "respetar", tiempo=8)
    assert bd.asignaciones(K).get((1, f), "L") == "L"  # regenerar todo el mes sí puede mover otras celdas (por eso existe «semana»)
    # novedad sobrevenida + regenerar solo esa semana: nada fuera de la semana cambia
    bd.guardar_novedad(Novedad(2, "Incapacidad", date(2026, 10, 14), date(2026, 10, 16)))
    antes = bd.asignaciones(K)
    flujo.generar(bd, *P, "semana", lunes=date(2026, 10, 12), tiempo=8)
    despues = bd.asignaciones(K)
    fuera = lambda d: not date(2026, 10, 12) <= d <= date(2026, 10, 18)
    assert {k: v for k, v in antes.items() if fuera(k[1])} == {k: v for k, v in despues.items() if fuera(k[1])}
    assert all((2, date(2026, 10, d)) not in despues for d in (14, 15, 16))


if __name__ == "__main__":
    for nombre, fn in list(globals().items()):
        if nombre.startswith("test_"):
            fn()
    print("OK")
