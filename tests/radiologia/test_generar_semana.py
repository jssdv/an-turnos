"""Pruebas mínimas. Ejecutar: .venv/bin/python tests/test_generar_semana.py"""
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ))

from an.radiologia import analisis as a
from an.radiologia import generar_semana as g
from an.radiologia.generar_semana import DIA_COMPLETO_DEF, SLOTS, linea_tiempo

# --- regla de disponibilidad -------------------------------------------------
# Sábado Abdomen: solo tarde presencial; el mismo médico de disp cubre antes y después.
assert linea_tiempo([(13, 19, False, "Salom")], [], "Salom", "Salom", DIA_COMPLETO_DEF) == [
    ("12:00 am - 1:00 pm", "Salom", "Disponibilidad", "disp"),
    ("1:00 pm - 7:00 pm", "Salom", None, "turno"),
    ("7:00 pm - 11:59 pm", "Salom", "Disponibilidad", "disp"),
]
# Viernes Neuro: corrido presencial, franjas remotas vacías, disp desde las 7 pm.
filas = linea_tiempo([(7, 13, False, "Muñoz"), (13, 19, False, "Muñoz")], SLOTS["NEURO"],
                     "Ramirez", "Ramirez", DIA_COMPLETO_DEF)
assert filas[0] == ("12:00 am - 7:00 am", "Ramirez", "Disponibilidad", "disp")
assert ("7:00 am - 1:00 pm", None, "Remoto", "remoto") in filas
assert filas[-1] == ("7:00 pm - 11:59 pm", "Ramirez", "Disponibilidad", "disp")
# Día sin turnos: disp 24h; si cambia el médico de disp, se parte a las 7 am.
assert linea_tiempo([], [], "Torres", "Torres", DIA_COMPLETO_DEF) == [
    ("12:00 am - 11:59 pm (24h)", "Torres", "Disponibilidad", "disp")]
assert [f[0] for f in linea_tiempo([], [], "A", "B", DIA_COMPLETO_DEF)] == [
    "12:00 am - 7:00 am", "7:00 am - 11:59 pm"]

# --- con el cuadro real de octubre -------------------------------------------
ORIGEN = RAIZ / "ejemplos" / "radiologia" / "Cuadro Octubre 2026 (prueba con fechas).xlsx"
if ORIGEN.exists():
    tmp = Path(tempfile.mkdtemp())
    # Copia con fechas completas en la fila 3 (opción B), por si el original aún tiene números
    wb = openpyxl.load_workbook(ORIGEN, data_only=True)
    ws = wb["Cambios Octubre 2026"]
    f = datetime(2026, 10, 1)
    for c in range(6, ws.max_column + 1):
        v = ws.cell(3, c).value
        if isinstance(v, (int, float, datetime)) and not isinstance(v, bool):
            ws.cell(3, c).value, f = f, f + timedelta(1)
    wb.save(tmp / "mensual.xlsx")
    ws = g.cargar_mensual(tmp / "mensual.xlsx")

    assert g.lunes_disponibles(ws) == [date(2026, 10, d) for d in (5, 12, 19, 26)]
    lunes = date(2026, 10, 12)
    libro, resumen, datos = g.construir(ws, lunes)
    assert resumen["festivos"] == [lunes]

    # El mensual marca OJO (36 h en 3 días) a Silva el 13 y el 14
    assert [(x[0].day, x[1][:5]) for x in a.alertas(ws, lunes) if "36 horas" in x[2]] == [(13, "Silva"), (14, "Silva")]

    # Equidad: el total del mes es el que calcula el cuadro
    meneses = next(e for e in a.equidad(ws, lunes) if e["medico"].startswith("Meneses"))
    assert meneses["mes"] == 151 and meneses["jornada"] == 42

    # Agenda: Ramirez tiene TN8R (2-10 pm remoto) martes a jueves
    ev = a.agenda_por_medico(datos)["Ramirez Mejía, Nestor Julián"]
    assert [e[0].day for e in ev if e[2] == "Lectura Neuro (remoto)"] == [13, 14, 15]
    assert a.ics("Ramirez", ev).count("BEGIN:VEVENT") == len(ev)

    # Cambios: una celda editada = 1 cambio, marcado en amarillo
    viejo, _, _ = g.construir(ws, lunes)
    viejo["Ecografía - ECO"]["G3"] = "Otro, Medico"
    nuevo, _, _ = g.construir(ws, lunes)
    cambios = a.comparar_libros(viejo, nuevo, lunes)
    assert len(cambios) == 1 and "Otro, Medico" in cambios[0], cambios
    assert nuevo["Ecografía - ECO"]["G3"].fill.fgColor.rgb == "FFFFF59D"

    # Revisión: el 5/10 Laverde tiene CJ#/TURG (Ciudad Jardín 1-5 pm se cruza con Urgencias 2-10 pm)
    _, _, datos5 = g.construir(ws, date(2026, 10, 5))
    cruces = [r for r in a.revisar_cuadro(datos5) if r[2].startswith("Se cruzan")]
    assert [(r[0].day, r[1][:7]) for r in cruces] == [(5, "Laverde")], cruces

    # Varios meses: el mismo cuadro dos veces suma el doble
    dos = next(f for f in a.equidad_meses([("oct", ws), ("oct", ws)]) if f["medico"].startswith("Meneses"))
    assert dos["mes"] == 302 and len(dos["meses"]) == 2

    # Imprimible: una sección por hoja y los nombres adentro
    html = a.html_planilla(libro, lunes, [("Ecografía - ECO", "Ecografía", 3), ("Lectura - Neuro", "Neuro", 3)])
    assert html.count("<section>") == 2 and "Santacruz Carmen, Sebastián" in html

    # Versiones: SEMANA..., luego v2
    assert g.ruta_libre(tmp, lunes).name == "SEMANA 12 AL 18 DE OCTUBRE.xlsx"
    libro.save(g.ruta_libre(tmp, lunes))
    assert g.ruta_libre(tmp, lunes).name == "SEMANA 12 AL 18 DE OCTUBRE v2.xlsx"
    assert a.version_anterior(tmp, lunes).name == "SEMANA 12 AL 18 DE OCTUBRE.xlsx"
print("OK")
