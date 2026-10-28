import os, sys; sys.path.insert(0, os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
"""Pruebas de los exportes en Excel (necesitan AN_PG_URL). Imprime OK."""
import psycopg
from io import BytesIO

from fastapi import Depends
from fixtura_api import ORIGEN, entorno

from an.web.deps import contexto_area_lectura

H = {"Origin": ORIGEN}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

from an.web import conexion


def test_tabla_exporte_rls():
    with entorno() as e:
        if e is None:
            return
        pid = e.su.execute("SELECT id FROM periodo WHERE area_id = %s", (e.ar["A"],)).fetchone()[0]
        e.su.execute("INSERT INTO exporte (clinica_id, area_id, periodo_id, tipo, numero, nombre_archivo, contenido) VALUES (%s, %s, %s, 'cuadro', 1, 'x.xlsx', %s)",
                     (e.cl["A"], e.ar["A"], pid, b"abc"))
        e.su.commit()

        def contar(n, sql="SELECT count(*) FROM exporte"):
            with e.pool.connection() as conn:
                conexion.alcance(conn, e.cl[n], e.ar[n])
                f = conn.execute(sql).fetchone()
                return f[0] if f else None
        assert contar("A") == 1 and contar("B") == 0
        for sql in ("UPDATE exporte SET numero = 2", "DELETE FROM exporte"):
            try:
                contar("A", sql)
                assert False, sql
            except psycopg.errors.InsufficientPrivilege:
                pass


def sesion(e, correo):
    cli = e.cliente()
    e.entrar(cli, correo)
    return cli


def pid(e, n="A"):
    return e.su.execute("SELECT id FROM periodo WHERE area_id = %s", (e.ar[n],)).fetchone()[0]


def ruta(e, n="A", extra="", p=None):
    return f"/api/areas/{e.ar[n]}/periodos/{p or pid(e, n)}/exportes{extra}"


def auditoria(e, evento="exportacion"):
    return e.su.execute("SELECT metodo, ruta, codigo, usuario_id FROM auditoria WHERE evento = %s ORDER BY id", (evento,)).fetchall()


def test_deps_lectura():
    with entorno() as e:
        if e is None:
            return

        @e.app.post("/api/areas/{area}/ping_lectura")
        def ping_lectura(ctx=Depends(contexto_area_lectura, scope="function")):
            return {"ok": 1}
        url = f"/api/areas/{e.ar['A']}/ping_lectura"
        lector = sesion(e, "lector@a.co")
        assert lector.post(url, headers=H).status_code == 200
        assert lector.post(f"/api/areas/{e.ar['A']}/ping", headers=H).status_code == 403
        assert sesion(e, "editor@b.co").post(url, headers=H).status_code == 404


def test_crear_cuadro_y_version():
    with entorno() as e:
        if e is None:
            return
        cli = sesion(e, "lector@a.co")
        r = cli.post(ruta(e), headers=H, json={"tipo": "cuadro"})
        assert r.status_code == 201, r.text
        d = r.json()
        assert (d["tipo"], d["numero"], d["nombre_archivo"], d["creado_por"]) == ("cuadro", 1, "CUADRO OCTUBRE 2026.xlsx", "lector@a.co")
        assert d["bytes"] > 0
        assert cli.post(ruta(e), headers=H, json={"tipo": "cuadro"}).json()["nombre_archivo"] == "CUADRO OCTUBRE 2026 v2.xlsx"
        d3 = cli.post(ruta(e), headers=H, json={"tipo": "cuadro"}).json()
        assert (d3["numero"], d3["nombre_archivo"]) == (3, "CUADRO OCTUBRE 2026 v3.xlsx")
        assert cli.post(ruta(e), headers=H, json={"tipo": "puestos"}).json()["nombre_archivo"] == "PUESTOS OCTUBRE 2026.xlsx"
        d = cli.post(ruta(e), headers=H, json={"tipo": "horas"}).json()
        assert (d["numero"], d["nombre_archivo"]) == (1, "HORAS Y RECARGOS OCTUBRE 2026.xlsx")


def test_errores_de_crear():
    with entorno() as e:
        if e is None:
            return
        cli = sesion(e, "editor@a.co")
        for t in ("semana", "x"):
            assert cli.post(ruta(e), headers=H, json={"tipo": t}).status_code == 422
        vacio = cli.post(f"/api/areas/{e.ar['A']}/periodos", headers=H, json={"inicio": "2026-11-01", "fin": "2026-11-30"}).json()
        r = cli.post(ruta(e, p=vacio["id"]), headers=H, json={"tipo": "cuadro"})
        assert (r.status_code, r.json()["detalle"]) == (409, "No hay cuadro para exportar.")
        assert cli.post(ruta(e, p=999999), headers=H, json={"tipo": "cuadro"}).status_code == 404


def test_lista_y_auditoria():
    with entorno() as e:
        if e is None:
            return
        cli = sesion(e, "lector@a.co")
        for t in ("cuadro", "cuadro", "horas"):
            cli.post(ruta(e), headers=H, json={"tipo": t})
        lista = cli.get(ruta(e)).json()
        assert [(x["tipo"], x["numero"]) for x in lista] == [("horas", 1), ("cuadro", 2), ("cuadro", 1)]
        assert "contenido" not in lista[0]
        assert [(m, c) for m, _, c, _ in auditoria(e)] == [("POST", 201)] * 3


def test_descargar_y_auditar():
    with entorno() as e:
        if e is None:
            return
        cli = sesion(e, "lector@a.co")
        x = cli.post(ruta(e), headers=H, json={"tipo": "cuadro"}).json()
        r = cli.get(ruta(e, extra=f"/{x['id']}/descargar"))
        assert r.status_code == 200 and r.headers["content-type"] == XLSX
        assert r.headers["content-disposition"] == ("attachment; filename=\"CUADRO OCTUBRE 2026.xlsx\"; filename*=UTF-8''CUADRO%20OCTUBRE%202026.xlsx")
        from openpyxl import load_workbook
        assert load_workbook(BytesIO(r.content)).worksheets[0]["A1"].value == "Colaborador"
        a = auditoria(e)
        assert [(m, c) for m, _, c, _ in a] == [("POST", 201), ("GET", 200)] and a[1][1].endswith("/descargar") and a[1][3] == e.us["lector@a.co"]


def test_otra_clinica_404():
    with entorno() as e:
        if e is None:
            return
        x = sesion(e, "lector@a.co").post(ruta(e), headers=H, json={"tipo": "cuadro"}).json()
        b = sesion(e, "editor@b.co")
        pb = pid(e, "B")
        for url in (ruta(e), ruta(e, extra=f"/{x['id']}/descargar"), ruta(e, "B", f"/{x['id']}/descargar"), ruta(e, "B", p=pid(e)), ruta(e, "B", p=pb, extra=f"/{x['id']}/descargar")):
            assert b.get(url).status_code == 404, url
        assert b.post(ruta(e), headers=H, json={"tipo": "cuadro"}).status_code == 404
        assert all(c != 200 for _, _, c, u in auditoria(e) if u == e.us["editor@b.co"])


def test_exporte_de_otro_periodo_404():
    with entorno() as e:
        if e is None:
            return
        cli = sesion(e, "editor@a.co")
        x = cli.post(ruta(e), headers=H, json={"tipo": "cuadro"}).json()
        otro = cli.post(f"/api/areas/{e.ar['A']}/periodos", headers=H, json={"inicio": "2026-11-01", "fin": "2026-11-30"}).json()["id"]
        assert cli.get(ruta(e, p=otro, extra=f"/{x['id']}/descargar")).status_code == 404


def celdas(libro):
    """Todo lo que ve quien abre el libro: hojas, dimensiones, valor y estilo de cada celda, anchos y panel fijo."""
    out = []
    for ws in libro.worksheets:
        out.append((ws.title, ws.dimensions, ws.freeze_panes, sorted((k, v.width) for k, v in ws.column_dimensions.items())))
        for fila in ws.iter_rows():
            out += [(c.coordinate, c.value, c.fill.fgColor.rgb, c.font.bold, c.font.color.rgb if c.font.color else None, c.alignment.horizontal) for c in fila]
    return out


def test_identico_a_qt():
    import tempfile
    from pathlib import Path

    from openpyxl import load_workbook

    from an.turnos import exportar, flujo
    from an.turnos.db import BD
    from an.turnos.pg import Conexion
    with entorno() as e:
        if e is None:
            return
        cli = sesion(e, "lector@a.co")
        k = "2026-10-01/2026-10-31"
        inicio, fin = flujo.desde_clave(k)
        fechas, txt = flujo.fechas(inicio, fin), flujo.etiqueta(inicio, fin)
        with e.pool.connection() as conn:  # como an_app: el superusuario se salta RLS y vería también la clínica B
            conexion.alcance(conn, e.cl["A"], e.ar["A"])
            bd = BD.postgres(Conexion(conn, confirmar=False))
            ctx, cuadro = bd.contexto(2026), bd.asignaciones(k)
            with tempfile.TemporaryDirectory() as t:
                esperado = {
                    "cuadro": exportar.exportar_cuadro(t, txt, ctx, cuadro, bd.asignaciones(k, "original"), fechas, bd.diferencias(k), bd.asignacion_puesto(k), bd.catalogo_puestos(),
                                                       flujo.detalle_horas(bd, inicio, fin)),
                    "puestos": exportar.exportar_puestos(t, txt, ctx, cuadro, fechas, bd.asignacion_puesto(k), bd.capacitacion(k), bd.catalogo_puestos()),
                    "horas": exportar.exportar_horas(t, txt, ctx, bd.tarifas(), cuadro, fechas, int(bd.param("ventana_semanas", 3)))}
                esperado = {tipo: celdas(load_workbook(r)) for tipo, r in esperado.items()}
        for tipo, celdas_qt in esperado.items():
            x = cli.post(ruta(e), headers=H, json={"tipo": tipo}).json()
            web = celdas(load_workbook(BytesIO(cli.get(ruta(e, extra=f"/{x['id']}/descargar")).content)))
            assert len(web) > 50, tipo
            assert web == celdas_qt, (tipo, len(web), len(celdas_qt), [(a, b) for a, b in zip(web, celdas_qt) if a != b][:3])


if __name__ == "__main__":
    import inspect
    for n, f in list(globals().items()):
        if n.startswith("test_") and inspect.isfunction(f):
            f()
    print("OK")


def test_tope_de_exportes_por_periodo_y_dia():
    from an.web.rutas_exportes import MAX_EXPORTES_DIA
    with entorno() as e:
        if e is None:
            return
        cli = sesion(e, "lector@a.co")
        for _ in range(MAX_EXPORTES_DIA):
            assert cli.post(ruta(e), headers=H, json={"tipo": "horas"}).status_code == 201
        r = cli.post(ruta(e), headers=H, json={"tipo": "horas"})
        assert r.status_code == 429 and "exportes" in r.json()["detalle"], r.text
        assert e.su.execute("SELECT count(*) FROM exporte").fetchone()[0] == MAX_EXPORTES_DIA
