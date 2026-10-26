import os, sys; sys.path.insert(0, os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
"""API de Radiología (necesitan AN_PG_URL). Imprime OK."""
import io
import zipfile

from fixtura_api import ORIGEN, entorno
from fixtura_radiologia import MEDICO, cuadro_minimo

from an.web.rutas_radiologia import MAX_ARCHIVO

CABECERAS = {"Origin": ORIGEN, "Content-Type": "application/octet-stream"}


def con_area_radiologia(e):
    """Un área `radiologia` por clínica; el editor/lector de A y el editor de B reciben permiso."""
    e.rad = {}
    for n in ("A", "B"):
        e.rad[n] = e.su.execute("INSERT INTO area (nombre, tipo, clinica_id) VALUES ('Radiología', 'radiologia', %s) RETURNING id", (e.cl[n],)).fetchone()[0]
    for correo, n, perm in (("editor@a.co", "A", "editar"), ("lector@a.co", "A", "ver"), ("editor@b.co", "B", "editar")):
        e.su.execute("INSERT INTO usuario_area (usuario_id, area_id, permiso, clinica_id) VALUES (%s, %s, %s, %s)", (e.us[correo], e.rad[n], perm, e.cl[n]))
    e.su.commit()


def subir(cli, area, datos, nombre="cuadro.xlsx"):
    return cli.post(f"/api/areas/{area}/archivos?nombre={nombre}", content=datos, headers=CABECERAS)


def test_subir_y_listar():
    with entorno() as e:
        if e is None:
            return
        con_area_radiologia(e)
        cli = e.cliente()
        e.entrar(cli, "editor@a.co")
        datos = cuadro_minimo()
        r = subir(cli, e.rad["A"], datos)
        assert r.status_code == 201 and r.json()["tamano"] == len(datos) and r.json()["nombre"] == "cuadro.xlsx", r.text
        lista = cli.get(f"/api/areas/{e.rad['A']}/archivos").json()
        assert [f["id"] for f in lista] == [r.json()["id"]] and "contenido" not in lista[0]
        assert cli.delete(f"/api/areas/{e.rad['A']}/archivos/{r.json()['id']}", headers={"Origin": ORIGEN}).status_code == 204
        assert cli.get(f"/api/areas/{e.rad['A']}/archivos").json() == []


def test_rechaza_lo_que_no_es_xlsx():
    with entorno() as e:
        if e is None:
            return
        con_area_radiologia(e)
        cli = e.cliente()
        e.entrar(cli, "editor@a.co")
        r = subir(cli, e.rad["A"], cuadro_minimo(), "cuadro.txt")
        assert r.status_code == 422 and r.json()["detalle"] == "Sube un archivo .xlsx."
        r = subir(cli, e.rad["A"], b"hola")
        assert r.status_code == 422 and r.json()["detalle"] == "El archivo no es un Excel .xlsx válido (¿está protegido con contraseña?)."


def test_rechaza_pasa_de_10_mb():
    with entorno() as e:
        if e is None:
            return
        con_area_radiologia(e)
        cli = e.cliente()
        e.entrar(cli, "editor@a.co")
        r = subir(cli, e.rad["A"], b"x" * (MAX_ARCHIVO + 1))
        assert r.status_code == 413 and r.json()["detalle"] == "El archivo pasa de 10 MB."
        r = subir(cli, e.rad["A"], b"x" * (MAX_ARCHIVO + 5000))
        assert r.status_code == 413 and r.json()["detalle"] == "El archivo pasa de 10 MB."
        assert subir(cli, e.rad["A"], b"x" * MAX_ARCHIVO).status_code == 422  # llega a la validación del contenido


def test_rechaza_zip_bomba():
    with entorno() as e:
        if e is None:
            return
        con_area_radiologia(e)
        cli = e.cliente()
        e.entrar(cli, "editor@a.co")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("relleno.bin", bytes(101 * 1024 * 1024))
        assert len(buf.getvalue()) < MAX_ARCHIVO
        r = subir(cli, e.rad["A"], buf.getvalue())
        assert r.status_code == 422 and r.json()["detalle"] == "El archivo es demasiado grande al abrirlo."


def test_cuadro_sin_excel_y_sin_semana():
    with entorno() as e:
        if e is None:
            return
        con_area_radiologia(e)
        cli = e.cliente()
        e.entrar(cli, "editor@a.co")
        r = subir(cli, e.rad["A"], cuadro_minimo(fila3=False))
        assert r.status_code == 422 and r.json()["detalle"] == (
            "La fila 3 de 'Cambios Octubre 2026' no tiene fechas completas.\nEscribe la primera fecha (ej. 01/10/2026), "
            "deja =+celda_anterior+1 en las demás, usa formato 'd' y guarda en Excel.")
        r = subir(cli, e.rad["A"], cuadro_minimo(dias=3))
        assert r.status_code == 422 and r.json()["detalle"] == "El cuadro no tiene ninguna semana completa (lunes a domingo)."


def test_permisos_y_aislamiento():
    with entorno() as e:
        if e is None:
            return
        con_area_radiologia(e)
        lector, editor, otro = e.cliente(), e.cliente(), e.cliente()
        e.entrar(lector, "lector@a.co")
        e.entrar(editor, "editor@a.co")
        e.entrar(otro, "editor@b.co")
        id_ = subir(editor, e.rad["A"], cuadro_minimo()).json()["id"]
        assert subir(lector, e.rad["A"], cuadro_minimo()).status_code == 403
        assert lector.delete(f"/api/areas/{e.rad['A']}/archivos/{id_}", headers={"Origin": ORIGEN}).status_code == 403
        assert len(lector.get(f"/api/areas/{e.rad['A']}/archivos").json()) == 1
        assert otro.get(f"/api/areas/{e.rad['A']}/archivos").status_code == 404  # otra clínica
        assert otro.delete(f"/api/areas/{e.rad['B']}/archivos/{id_}", headers={"Origin": ORIGEN}).status_code == 404  # archivo ajeno
        assert editor.get(f"/api/areas/{e.ar['A']}/archivos").status_code == 404  # área de turnos
        assert editor.get(f"/api/areas/{e.rad['B']}/archivos").status_code == 404  # sin permiso en esa área
        fila = e.su.execute("SELECT ruta FROM auditoria WHERE metodo = 'POST' AND ruta LIKE '%archivos%' ORDER BY id DESC LIMIT 1").fetchone()
        assert fila and "cuadro" not in fila["ruta"] and "xlsx" not in fila["ruta"], fila["ruta"]  # el nombre del archivo no llega a la auditoría


def test_semana_por_defecto():
    from datetime import date
    from an.radiologia.analisis import semana_por_defecto
    lunes = [date(2026, 10, d) for d in (5, 12, 19)]
    assert semana_por_defecto(lunes, date(2026, 10, 14)) == date(2026, 10, 12)  # la de esa semana
    assert semana_por_defecto(lunes, date(2026, 10, 3)) == date(2026, 10, 5)
    assert semana_por_defecto(lunes, date(2027, 1, 1)) == date(2026, 10, 19)


def con_cuadro(e, usuario="lector@a.co"):
    """(cliente lector, id del archivo subido por el editor, área de Radiología)."""
    con_area_radiologia(e)
    editor, cli = e.cliente(), e.cliente()
    e.entrar(editor, "editor@a.co")
    e.entrar(cli, usuario)
    return cli, subir(editor, e.rad["A"], cuadro_minimo()).json()["id"], e.rad["A"]


def test_semanas_y_analisis():
    with entorno() as e:
        if e is None:
            return
        cli, id_, area = con_cuadro(e)
        s = cli.get(f"/api/areas/{area}/archivos/{id_}/semanas").json()
        assert s["hoja"] == "Cambios Octubre 2026"
        assert [(x["lunes"], x["etiqueta"]) for x in s["semanas"]] == [
            ("2026-10-05", "5 – 11 oct 2026"), ("2026-10-12", "12 – 18 oct 2026"), ("2026-10-19", "19 – 25 oct 2026")]
        assert s["defecto"] in {x["lunes"] for x in s["semanas"]}
        d = cli.get(f"/api/areas/{area}/archivos/{id_}/semanas/2026-10-05").json()
        r = d["resumen"]
        assert r["semana"] == "Semana del 5 de octubre al 11 de octubre"
        assert [(t["titulo"], t["valor"]) for t in r["tarjetas"]] == [
            ("Franjas sin médico", 104), ("Alertas de descanso", 0), ("Problemas en el cuadro", 0),
            ("Cambios desde la última versión", "—"), ("Médicos con turnos", 1), ("Festivos", 0)]
        abd = next(f for f in r["cobertura"] if f["servicio"] == "Lectura Abdomen")
        assert abd["dias"][0] == {"n": 1, "faltan": ["12:00 am - 7:00 am", "1:00 pm - 11:59 pm", "1:00 pm - 7:00 pm"]}
        assert r["aviso"] == "" and d["alertas"]["descanso"] == [] and d["alertas"]["revision"] == []
        assert len(d["alertas"]["pendientes"]) == 104
        assert d["equidad"][0]["medico"] == MEDICO and d["equidad"][0]["semana"] == 6
        assert list(d["agendas"]) == [MEDICO] and d["agendas"][MEDICO].strip()
        assert d["cambios"] == {"version_anterior": None, "lista": []}
        for malo in ("2026-10-06", "2026-11-02", "no-es-fecha"):  # martes, fuera del cuadro, basura
            x = cli.get(f"/api/areas/{area}/archivos/{id_}/semanas/{malo}")
            assert x.status_code == 422 and x.json()["detalle"] == "La fecha debe ser un lunes del cuadro."
        assert cli.get(f"/api/areas/{area}/archivos/9999/semanas").status_code == 404


def test_previa():
    with entorno() as e:
        if e is None:
            return
        cli, id_, area = con_cuadro(e)
        hojas = cli.get(f"/api/areas/{area}/archivos/{id_}/semanas/2026-10-05/previa").json()
        assert [h["titulo"] for h in hojas] == ["Ecografía", "Neuro", "MSK", "Abdomen", "Tórax y Corazón", "Mamografía y Densito"]  # sin «Cambios»: no hay versión anterior
        abd = next(h for h in hojas if h["hoja"] == "Lectura - Abdomen")
        assert any(c["texto"] == MEDICO for fila in abd["filas"] for c in fila) and abd["fusiones"]


def test_equidad_meses():
    with entorno() as e:
        if e is None:
            return
        cli, id_, area = con_cuadro(e)
        editor = e.cliente()
        e.entrar(editor, "editor@a.co")
        id2 = subir(editor, area, cuadro_minimo()).json()["id"]
        r = cli.get(f"/api/areas/{area}/equidad-meses?archivos={id_},{id2}").json()
        assert r["meses"] == ["Octubre 2026", "Octubre 2026"] and r["fallidos"] == []
        assert r["filas"][0]["medico"] == MEDICO and r["filas"][0]["mes"] == 20  # 10 + 10
        assert cli.get(f"/api/areas/{area}/equidad-meses?archivos=9999").status_code == 404
        assert cli.get(f"/api/areas/{area}/equidad-meses?archivos=a").status_code == 422
        assert cli.get(f"/api/areas/{area}/equidad-meses?archivos=" + ",".join(str(i) for i in range(1, 14))).status_code == 422


def test_agendas_ics_y_zip():
    import urllib.parse
    from an.radiologia import analisis, generar_semana as gs
    with entorno() as e:
        if e is None:
            return
        cli, id_, area = con_cuadro(e)
        base = f"/api/areas/{area}/archivos/{id_}/semanas/2026-10-05"
        r = cli.get(f"{base}/agenda.ics?medico={urllib.parse.quote(MEDICO)}")
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/calendar") and r.text.startswith("BEGIN:VCALENDAR")
        assert r.text.count("BEGIN:VEVENT") >= 1
        assert 'filename="Prueba Uno, Ana.ics"' in r.headers["content-disposition"]
        assert cli.get(f"{base}/agenda.ics?medico=Nadie").status_code == 404
        z = zipfile.ZipFile(io.BytesIO(cli.get(f"{base}/agendas.zip").content))
        assert sorted(z.namelist()) == [f"{MEDICO}.ics", f"{MEDICO}.txt"]
        assert z.read(f"{MEDICO}.ics").decode().count("BEGIN:VEVENT") == r.text.count("BEGIN:VEVENT")
    assert analisis.nombre_archivo('A/B:"C"') == "ABC"


H = {"Origin": ORIGEN}


def exportar(cli, area, id_, lunes):
    return cli.post(f"/api/areas/{area}/archivos/{id_}/semanas/{lunes}/exportes", headers=H)


def bajar(cli, area, x):
    return cli.get(f"/api/areas/{area}/periodos/{x['periodo_id']}/exportes/{x['id']}/descargar")


def test_exporte_numera_sin_sobrescribir_y_audita():
    import openpyxl
    from an.radiologia import generar_semana as gs
    with entorno() as e:
        if e is None:
            return
        cli, id_, area = con_cuadro(e)
        a = exportar(cli, area, id_, "2026-10-19")
        b = exportar(cli, area, id_, "2026-10-19")
        assert (a.status_code, a.json()["nombre_archivo"], a.json()["numero"]) == (201, "SEMANA 19 AL 25 DE OCTUBRE.xlsx", 1), a.text
        assert (b.json()["nombre_archivo"], b.json()["numero"]) == ("SEMANA 19 AL 25 DE OCTUBRE v2.xlsx", 2)
        otra = exportar(cli, area, id_, "2026-10-12").json()
        assert (otra["nombre_archivo"], otra["numero"]) == ("SEMANA 12 AL 18 DE OCTUBRE.xlsx", 1)
        lista = cli.get(f"/api/areas/{area}/archivos/{id_}/semanas/2026-10-19/exportes").json()
        assert [x["numero"] for x in lista] == [2, 1]
        r = bajar(cli, area, a.json())
        assert r.status_code == 200 and openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames == openpyxl.load_workbook(gs.PLANTILLA).sheetnames
        eventos = [(f["metodo"], f["codigo"]) for f in e.su.execute("SELECT metodo, codigo FROM auditoria WHERE evento = 'exportacion' ORDER BY id")]
        assert eventos == [("POST", 201)] * 3 + [("GET", 200)]
        assert e.su.execute("SELECT count(*) FROM exporte WHERE tipo = 'semana'").fetchone()[0] == 3


def test_cambios_contra_la_version_guardada():
    with entorno() as e:
        if e is None:
            return
        cli, id_, area = con_cuadro(e)
        editor = e.cliente()
        e.entrar(editor, "editor@a.co")
        base = f"/api/areas/{area}/archivos"
        assert cli.get(f"{base}/{id_}/semanas/2026-10-19").json()["cambios"] == {"version_anterior": None, "lista": []}
        exportar(cli, area, id_, "2026-10-19")
        igual = cli.get(f"{base}/{id_}/semanas/2026-10-19").json()
        assert igual["cambios"] == {"version_anterior": "SEMANA 19 AL 25 DE OCTUBRE.xlsx", "lista": []}
        assert igual["resumen"]["tarjetas"][3]["valor"] == 0
        # el cuadro corregido: el lunes 19 pasa de MA (mañana) a TA (tarde)
        nuevo = subir(editor, area, cuadro_minimo(codigos={15: "TA"})).json()["id"]
        d = cli.get(f"{base}/{nuevo}/semanas/2026-10-19").json()
        assert d["cambios"]["version_anterior"] == "SEMANA 19 AL 25 DE OCTUBRE.xlsx" and len(d["cambios"]["lista"]) >= 1
        n2 = exportar(cli, area, nuevo, "2026-10-19").json()
        assert n2["nombre_archivo"].endswith(" v2.xlsx") and n2["cambios"] == len(d["cambios"]["lista"])
        import openpyxl
        guardado = openpyxl.load_workbook(io.BytesIO(bajar(cli, area, n2).content))
        assert guardado.sheetnames[0] == "Cambios"  # el libro guardado lleva la hoja «Cambios» cuando hay diferencias


def _preparado(origen, destino):
    """Copia con fechas completas en la fila 3 (como en tests/radiologia/test_generar_semana.py)."""
    from datetime import datetime, timedelta
    import openpyxl
    wb = openpyxl.load_workbook(origen, data_only=True)
    ws = wb["Cambios Octubre 2026"]
    f = datetime(2026, 10, 1)
    for c in range(6, ws.max_column + 1):
        v = ws.cell(3, c).value
        if isinstance(v, (int, float, datetime)) and not isinstance(v, bool):
            ws.cell(3, c).value, f = f, f + timedelta(1)
    wb.save(destino)


def _diferencias(a, b):
    """Celdas distintas entre dos libros (valor, relleno, negrita, combinadas); solo coordenadas, nunca valores."""
    malas = []
    assert a.sheetnames == b.sheetnames
    for nombre in a.sheetnames:
        x, y = a[nombre], b[nombre]
        if sorted(map(str, x.merged_cells.ranges)) != sorted(map(str, y.merged_cells.ranges)):
            malas.append(f"{nombre}: combinadas")
        for fila in range(1, max(x.max_row, y.max_row) + 1):
            for col in range(1, max(x.max_column, y.max_column) + 1):
                p, q = x.cell(fila, col), y.cell(fila, col)
                if (p.value, p.fill.fgColor.rgb, p.font.b) != (q.value, q.fill.fgColor.rgb, q.font.b):
                    malas.append(f"{nombre}!{p.coordinate}")
    return malas


def test_libro_igual_a_qt():
    """Aceptación 7b: la semana del 19 al 25 de octubre sale igual (celdas) que la que genera el núcleo con las mismas entradas."""
    import tempfile
    from datetime import date
    from pathlib import Path
    import openpyxl
    from an.radiologia import generar_semana as gs
    ejemplo = Path(__file__).resolve().parents[2] / "ejemplos" / "radiologia" / "Cuadro Octubre 2026 (prueba con fechas).xlsx"
    with entorno() as e:
        if e is None:
            return
        tmp = Path(tempfile.mkdtemp())
        if ejemplo.exists():
            _preparado(ejemplo, tmp / "mensual.xlsx")
        else:  # sin los datos locales, el cuadro mínimo prueba el mismo camino
            (tmp / "mensual.xlsx").write_bytes(cuadro_minimo())
        con_area_radiologia(e)
        editor = e.cliente()
        e.entrar(editor, "editor@a.co")
        id_ = subir(editor, e.rad["A"], (tmp / "mensual.xlsx").read_bytes()).json()["id"]
        x = exportar(editor, e.rad["A"], id_, "2026-10-19")
        assert x.status_code == 201
        r = bajar(editor, e.rad["A"], x.json())
        assert r.status_code == 200
        gs.generar(tmp / "mensual.xlsx", date(2026, 10, 19), tmp / "qt.xlsx")
        malas = _diferencias(openpyxl.load_workbook(io.BytesIO(r.content)), openpyxl.load_workbook(tmp / "qt.xlsx"))
        assert not malas, malas[:20]
