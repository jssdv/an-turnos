"""Cuadro mensual mínimo en memoria (datos ficticios) para las pruebas de Radiología en la web."""
import io
from datetime import datetime, timedelta

import openpyxl

from an.nucleo.xlsx import poner_texto

MEDICO = "Prueba Uno, Ana"
HOJA = "Cambios Octubre 2026"


def cuadro_minimo(dias=22, fila3=True, hoja=HOJA, codigos=None, medico=MEDICO):
    """.xlsx con las fechas del 4 al 25 de octubre de 2026 (lunes 5, 12 y 19), un médico con `MA` (abdomen mañana)
    cada lunes y la hoja Nom-Hrs. `dias` recorta las columnas; `fila3=False` simula el cuadro guardado sin Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = hoja
    for c, t in zip(range(1, 6), ("", "ESP1", "ESP2", "TIPO", "MEDICO")):
        ws.cell(2, c).value = t or None
    for c, t in ((6 + dias + 0, "JORNADA"), (6 + dias + 1, "TOTAL")):
        ws.cell(2, c).value = t
    f = datetime(2026, 10, 4)
    for k in range(dias):
        if fila3:
            ws.cell(3, 6 + k).value = f + timedelta(k)
    ws.cell(4, 3).value = "ABDOMEN"
    poner_texto(ws.cell(4, 5), medico)  # como texto: un nombre que empieza por «=» llega así desde Excel
    for k in range(dias):
        if (f + timedelta(k)).weekday() == 0:
            ws.cell(4, 6 + k).value = (codigos or {}).get(k, "MA")
    ws.cell(4, 6 + dias).value = 42
    ws.cell(4, 6 + dias + 1).value = 10
    nh = wb.create_sheet("Nom-Hrs")
    nh.append(["", "", "COD", "HRS"])
    nh.append(["", "", "MA", 6])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cuadro_con_celda_lejana(fila=1048576):
    """El cuadro mínimo más una celda en `E<fila>`: pesa unos KB pero openpyxl cree que la hoja tiene un millón de filas."""
    wb = openpyxl.load_workbook(io.BytesIO(cuadro_minimo()))
    wb[HOJA].cell(fila, 5).value = "x"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
