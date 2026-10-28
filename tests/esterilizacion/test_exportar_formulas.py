import os, sys; sys.path.insert(0, os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")))
"""Un nombre que empieza por «=» sale como texto en el Excel exportado, no como fórmula."""
from datetime import date, timedelta
from types import SimpleNamespace as NS

from openpyxl import Workbook

from an.turnos import exportar, recargos


def test_nombre_hostil_sale_como_texto():
    hostil = '=HYPERLINK("https://x","Ver")'
    ctx = NS(colabs=[NS(id=1, nombre=hostil, en_solver=True)], turnos={}, tipo_dia=lambda f: "DL", novedad=lambda i, f: None, demandas=[], festivos=set())
    ws = Workbook().active
    exportar._hoja_cuadro(ws, ctx, {}, [date(2026, 10, 5) + timedelta(i) for i in range(3)])
    assert ws["A2"].value == hostil and ws["A2"].data_type == "s"


def test_cuadro_con_horas_y_contadores():
    import tempfile
    import openpyxl
    from an.turnos import db, flujo
    from an.turnos.dominio import Colaborador
    bd = db.BD()
    for i in range(2):
        bd.guardar_colaborador(Colaborador(0, f"P{i}"))
    ctx = bd.contexto(2026)
    a, b = ctx.colabs[0], ctx.colabs[1]
    ini = date(2026, 10, 5)
    fechas = [ini + timedelta(i) for i in range(7)]
    cuadro = {(a.id, fechas[0]): "M", (a.id, fechas[1]): "M", (a.id, fechas[2]): "NC", (b.id, fechas[0]): "T"}
    detalle = {p.id: recargos.reporte(ctx, bd.tarifas(), p, cuadro, fechas, 3) for p in ctx.colabs}
    with tempfile.TemporaryDirectory() as tmp:
        ruta = exportar.exportar_cuadro(tmp, "Prueba", ctx, cuadro, cuadro, fechas, [], detalle=detalle)
        wb = openpyxl.load_workbook(ruta)
    ws = wb["Cuadro"]
    enc = [c.value for c in ws[1]]
    assert enc[8:12] == ["Σ sem", "Horas", "Contratadas", "De más"] and "NC" in enc and "Horas equiv. recargos" in enc
    assert "Horas nocturnas" in enc and "Horas dom/fest" in enc
    assert ws["I2"].value.startswith("=SUMPRODUCT(COUNTIF(B2:H2,Turnos!")  # suma semanal tras el domingo
    assert ws["J2"].value.startswith("=SUMPRODUCT(COUNTIF($B2:$I2,Turnos!")
    assert ws.cell(2, enc.index("NC") + 1).value == '=COUNTIF($B2:$I2,"NC")'
    assert wb["Turnos"]["A2"].value in ctx.turnos
    textos = [c.value for c in ws["A"]]
    assert "Contador por turno" in textos and any(str(t).startswith("M ") for t in textos)
    assert ws["A3"].data_type == "s" and ws["B2"].data_type == "s"


if __name__ == "__main__":
    test_nombre_hostil_sale_como_texto()
    test_cuadro_con_horas_y_contadores()
    print("OK")
