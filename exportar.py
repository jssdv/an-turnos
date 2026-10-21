"""Salidas en Excel: cuadro mensual, planilla de puestos y reporte de horas y recargos. Nunca sobrescribe."""
from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import recargos
from dominio import DIAS
from puestos import LEYENDA, PUESTOS, etiqueta, por_persona, tiene_puesto

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre",
         "Noviembre", "Diciembre"]
ABREV_NOVEDAD = {"Vacaciones": "VAC", "Incapacidad": "INC", "Licencia": "LIC", "Libre solicitado": "LIB"}
ENCABEZADO = PatternFill("solid", fgColor="1F3864")
AMARILLO = PatternFill("solid", fgColor="FFFF00")
GRIS = PatternFill("solid", fgColor="D9D9D9")
FESTIVO = PatternFill("solid", fgColor="F8CBAD")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def ruta_libre(carpeta, base):
    """`base.xlsx`, luego `base v2.xlsx`, `base v3.xlsx`…"""
    carpeta, n = Path(carpeta), 1
    while (r := carpeta / (f"{base}.xlsx" if n == 1 else f"{base} v{n}.xlsx")).exists():
        n += 1
    return r


def _encabezado(ws, fila, textos, col=1):
    for i, t in enumerate(textos):
        c = ws.cell(fila, col + i, t)
        c.fill, c.font, c.alignment = ENCABEZADO, Font(bold=True, color="FFFFFF"), CENTRO


def _ancho(ws, primera, resto, desde=2):
    ws.column_dimensions["A"].width = primera
    for i in range(desde, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = resto


def _hoja_cuadro(ws, ctx, cuadro, fechas, comparar=None, puestos=None):
    cols = [p for p in ctx.colabs if p.en_solver]
    _encabezado(ws, 1, ["Colaborador"] + [f"{DIAS[f.weekday()]}\n{f.day}" for f in fechas] + ["Horas"])
    for i, f in enumerate(fechas, 2):
        if ctx.tipo_dia(f) == "DF":
            ws.cell(1, i).fill = PatternFill("solid", fgColor="C55A11")
    for r, p in enumerate(cols, 2):
        ws.cell(r, 1, p.nombre).font = Font(bold=True)
        horas = 0
        for i, f in enumerate(fechas, 2):
            cod = cuadro.get((p.id, f), "L")
            nov = ctx.novedad(p.id, f)
            texto = ABREV_NOVEDAD[nov.tipo] if nov and cod == "L" else cod
            if puestos is not None and cod in ctx.turnos and tiene_puesto(ctx.turnos[cod]) and (p.id, f) in puestos:
                texto = etiqueta(cod, puestos[p.id, f])
            c = ws.cell(r, i, texto)
            c.alignment = CENTRO
            if cod in ctx.turnos:
                c.fill = PatternFill("solid", fgColor=ctx.turnos[cod].color)
                horas += ctx.turnos[cod].horas
            elif nov:
                c.fill = GRIS
            if comparar is not None and comparar.get((p.id, f), "L") != cod:
                c.fill = AMARILLO
        ws.cell(r, len(fechas) + 2, horas)
    base = len(cols) + 3
    for k, (fr, nombre) in enumerate((("M", "Mañana"), ("T", "Tarde"), ("N", "Noche"))):
        ws.cell(base + k, 1, f"Cubiertos {nombre}").font = Font(bold=True)
        for i, f in enumerate(fechas, 2):
            n = sum(1 for p in cols if p.en_solver and (t := ctx.turnos.get(cuadro.get((p.id, f)))) and fr in t.cubre)
            req = ctx.cobertura[ctx.tipo_dia(f)][fr]
            c = ws.cell(base + k, i, n)
            c.alignment = CENTRO
            if n < req:
                c.fill = PatternFill("solid", fgColor="FF7C80")
    _ancho(ws, 22, 5)
    ws.column_dimensions[get_column_letter(len(fechas) + 2)].width = 8
    ws.freeze_panes = "B2"


def exportar_cuadro(carpeta, mes_txt, ctx, cuadro, original, fechas, diferencias, filas_puestos=()):
    """Hojas: Cuadro (de `cambios`, celdas modificadas en amarillo), Cuadro con puestos, Original y Cambios (diferencias)."""
    wb = Workbook()
    _hoja_cuadro(wb.active, ctx, cuadro, fechas, comparar=original)
    wb.active.title = "Cuadro"
    if filas_puestos:
        con = wb.create_sheet("Cuadro con puestos")
        _hoja_cuadro(con, ctx, cuadro, fechas, puestos=por_persona(filas_puestos))
        con.cell(con.max_row + 2, 1, LEYENDA)
    _hoja_cuadro(wb.create_sheet("Original"), ctx, original, fechas)
    ws = wb.create_sheet("Cambios")
    _encabezado(ws, 1, ["Colaborador", "Fecha", "Antes", "Después"])
    nombres = {p.id: p.nombre for p in ctx.colabs}
    for r, (cid, f, antes, ahora) in enumerate(diferencias, 2):
        for c, v in enumerate((nombres.get(cid, cid), f.strftime("%d/%m/%Y"), antes, ahora), 1):
            ws.cell(r, c, v)
    _ancho(ws, 26, 14)
    ruta = ruta_libre(carpeta, f"CUADRO {mes_txt.upper()}")
    wb.save(ruta)
    return ruta


def exportar_puestos(carpeta, mes_txt, ctx, cuadro, fechas, filas, capacitacion):
    """Una hoja por semana: por día, quién ocupa cada puesto en la mañana y la tarde, la noche y la capacitación."""
    wb = Workbook()
    wb.remove(wb.active)
    nombres = {p.id: p.nombre for p in ctx.colabs}
    cap = {}
    for f, cid in capacitacion:
        cap.setdefault(f, []).append(nombres[cid])
    cel = {(f, fr, p): f"{nombres[c]} (P{g})" for f, fr, p, c, g in filas if p}
    apoyo = {}
    for f, fr, p, c, g in filas:
        if not p:
            apoyo.setdefault((f, fr), []).append(nombres[c])
    puestos = [1, 2, 3, 4, 5]
    for lunes in sorted({f - timedelta(days=f.weekday()) for f in fechas}):
        ws = wb.create_sheet(f"Semana {lunes:%d-%m}")
        ws.cell(1, 1, "Día").fill = ENCABEZADO
        enc = ["Día"] + [f"M · {PUESTOS[p]}" for p in puestos] + ["M · Apoyo"] + [f"T · {PUESTOS[p]}" for p in puestos] \
            + ["T · Apoyo", "Noche", "Capacitación"]
        _encabezado(ws, 1, enc)
        for r, f in enumerate((lunes + timedelta(days=i) for i in range(7)), 2):
            c = ws.cell(r, 1, f"{DIAS[f.weekday()]} {f:%d/%m}")
            c.font = Font(bold=True)
            if ctx.tipo_dia(f) == "DF":
                c.fill = FESTIVO
            if f not in fechas:
                continue
            fila = [cel.get((f, "M", p), "") for p in puestos] + [", ".join(apoyo.get((f, "M"), []))] \
                + [cel.get((f, "T", p), "") for p in puestos] + [", ".join(apoyo.get((f, "T"), []))]
            noche = [nombres[cid] for (cid, d), cod in cuadro.items() if d == f and cod in ctx.turnos and "N" in ctx.turnos[cod].cubre]
            fila += [", ".join(sorted(noche)), ", ".join(cap.get(f, []))]
            for i, v in enumerate(fila, 2):
                ws.cell(r, i, v).alignment = CENTRO
        ws.cell(10, 1, "(P1)/(P2): grupo de pausa — el grupo 1 sale primero y el 2 cubre; se alternan cada día.")
        _ancho(ws, 14, 20)
        ws.freeze_panes = "B2"
    ruta = ruta_libre(carpeta, f"PUESTOS {mes_txt.upper()}")
    wb.save(ruta)
    return ruta


CATEGORIAS = [(t, n, f) for t in ("ordinaria", "extra") for f in ("hábil", "festivo") for n in ("diurna", "nocturna")]


def exportar_horas(carpeta, mes_txt, ctx, tarifas, cuadro, fechas, ventana_sem):
    wb = Workbook()
    ws = wb.active
    ws.title = "Horas y recargos"
    nombre_cat = lambda c: f"{'Extra' if c[0] == 'extra' else 'Ordinaria'} {c[1]} {'dom/fest' if c[2] == 'festivo' else 'hábil'}"
    _encabezado(ws, 1, ["Colaborador", "Contratadas", "Trabajadas", "Pendientes", "Adicionales"]
                + [nombre_cat(c) for c in CATEGORIAS] + ["Dom/fest trabajados", "Habitual", "Horas equivalentes de recargo"])
    sem = wb.create_sheet("Semanas")
    lunes = sorted({f - timedelta(days=f.weekday()) for f in fechas})
    _encabezado(sem, 1, ["Colaborador"] + [f"Sem {l:%d/%m}" for l in lunes])
    for r, p in enumerate((p for p in ctx.colabs if p.en_solver), 2):
        rep = recargos.reporte(ctx, tarifas, p, cuadro, fechas, ventana_sem)
        fila = [p.nombre] + [round(rep[k], 1) for k in ("contratadas", "trabajadas", "pendientes", "adicionales")] \
            + [round(rep["cat"].get(c, 0), 1) for c in CATEGORIAS] \
            + [rep["festivos_trabajados"], "Sí" if rep["habitual"] else "No", round(rep["equivalentes"], 1)]
        for c, v in enumerate(fila, 1):
            ws.cell(r, c, v)
        sem.cell(r, 1, p.nombre)
        for c, l in enumerate(lunes, 2):
            sem.cell(r, c, round(rep["semanas"].get(l, 0), 1))
    _ancho(ws, 24, 14)
    _ancho(sem, 24, 12)
    ws.freeze_panes = "B2"
    nota = wb.create_sheet("Régimen")
    f0 = fechas[0]
    filas = [("Régimen de horas", f"Art. 165 CST: promedio en ventanas de {ventana_sem} semanas contra la jornada contratada."),
             ("Ventana", f"Contada desde el primer día del mes; no cruza meses."),
             ("Pendientes / adicionales", "Déficit o exceso de la ventana frente a lo contratado (descontando novedades)."),
             ("Horas extra", "Se toman de las últimas horas trabajadas de la ventana."),
             ("Dominical/festivo", f"{recargos.tarifa(tarifas, 'dominical', f0)} %"),
             ("Nocturno", f"{recargos.tarifa(tarifas, 'nocturno', f0)} % desde las {recargos.tarifa(tarifas, 'noche_inicio', f0) // 60}:00"),
             ("Extra diurna / nocturna", f"{recargos.tarifa(tarifas, 'extra_diurna', f0)} % / {recargos.tarifa(tarifas, 'extra_nocturna', f0)} %"),
             ("Horas equivalentes", "Ordinarias: horas × recargo. Extras: horas × (1 + recargo). Los recargos se suman."),
             ("Aviso", "Cifras de apoyo para nómina; confirmar con RR. HH./jurídico.")]
    for r, (a, b) in enumerate(filas, 1):
        nota.cell(r, 1, a).font = Font(bold=True)
        nota.cell(r, 2, b)
    _ancho(nota, 28, 90)
    ruta = ruta_libre(carpeta, f"HORAS Y RECARGOS {mes_txt.upper()}")
    wb.save(ruta)
    return ruta
