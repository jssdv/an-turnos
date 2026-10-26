"""Análisis sobre el cuadro mensual y la planilla generada: alertas, equidad, agendas y cambios."""
import re
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl.styles import Font, PatternFill

from .generar_semana import HOJAS_LECTURA, fechas_mensual, hora, nombre_semana

DIAS = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
DIAS_LARGOS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
MESES_CORTOS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
MESES_LARGOS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
                "octubre", "noviembre", "diciembre"]
SERVICIOS = {"URG": "Ecografía Urgencias", "ECO1": "Ecógrafo 1 (Mixto)", "ECO2": "Ecógrafo 2 (Mama)",
             "ECO3": "Ecógrafo 3 (Abdomen)", "ECO4": "Ecógrafo 4 (MSK)", "ECO5": "Ecógrafo 5 (Procedimientos)",
             "ECO6": "Ecógrafo 6 (Mixto)", "PM": "Ecografía Pacific Mall", "CJ1": "Ecografía Ciudad Jardín 1",
             "CJ2": "Ecografía Ciudad Jardín 2", "NEURO": "Lectura Neuro", "MSK": "Lectura MSK",
             "ABD": "Lectura Abdomen", "TORAX": "Lectura Tórax", "CARDIO": "Lectura Corazón",
             "MAMA": "Lectura Mamografía", "DENSITO": "Lectura Densitometría",
             "PROCED": "Procedimientos"}
AMARILLO = PatternFill("solid", fgColor="FFFFF59D")
PESTANAS_PREVIA = {"Ecografía - ECO": "Ecografía", "Lectura - Neuro": "Neuro", "Lectura - MSK": "MSK",
                   "Lectura - Abdomen": "Abdomen", "Lectura - CARDIO-TORAX AMBULATO": "Tórax y Corazón",
                   "Lectura - Mamografia-Densitomet": "Mamografía y Densito", "Cambios": "Cambios"}
# Servicios en el orden de la tabla de cobertura
COBERTURA = ["URG", "ECO1", "ECO2", "ECO3", "ECO4", "ECO5", "PM", "CJ1", "CJ2",
             "NEURO", "MSK", "ABD", "TORAX", "CARDIO", "MAMA", "DENSITO"]


def fecha_corta(f):
    return f"{DIAS[f.weekday()]} {f:%d/%m}"


# ---------------------------------------------------------------- cuadro mensual
def filas_medicos(ws):
    """{nombre: fila principal} (filas sin 'Disp' en la columna D, hasta CONTINGENCIA)."""
    filas = {}
    for r in range(4, ws.max_row + 1):
        n = ws.cell(r, 5).value
        if isinstance(n, str) and n.startswith("CONTINGENCIA"):
            break
        if isinstance(n, str) and n.strip() and str(ws.cell(r, 4).value or "").strip().lower() != "disp":
            filas.setdefault(n.strip(), r)
    return filas


def horas_codigo(ws):
    """{código: horas} de la hoja Nom-Hrs (valores ya calculados por Excel)."""
    nh = ws.parent["Nom-Hrs"]
    return {str(c).strip(): (h if isinstance(h, (int, float)) else 0)
            for c, h in nh.iter_rows(min_row=2, min_col=3, max_col=4, values_only=True) if c}


def alertas(ws, lunes):
    """Reglas de descanso del mensual (filas "prueba turg/nurg/CCC") para la semana.

    Devuelve [(fecha, médico, mensaje)].
    """
    fechas, horas = fechas_mensual(ws), horas_codigo(ws)
    out = []
    for nombre, r in filas_medicos(ws).items():
        def cod(f):
            v = ws.cell(r, fechas[f]).value if f in fechas else None
            return v.strip() if isinstance(v, str) else ""

        for k in range(7):
            f = lunes + timedelta(k)
            hoy, man = cod(f), cod(f + timedelta(1))
            if hoy == "TURG" and man[:1] in ("M", "C"):
                out.append((f, nombre, f"TURG seguido de {man} al día siguiente"))
            if hoy == "NURG" and (f + timedelta(1)) in fechas and man != "LP":
                out.append((f, nombre, f"NURG sin libre posturno (sigue {man or 'vacío'})"))
            tres = [cod(f + timedelta(i)) for i in range(3)]
            if sum(horas.get(c, 0) for c in tres) >= 36:
                out.append((f, nombre, f"36 horas en 3 días ({' · '.join(tres)})"))
            # tarde larga (termina 9-10 pm) y al otro día entra temprano; TURG ya lo cubre la 1ª regla
            if hoy[:1] == "T" and hoy != "TURG" and horas.get(hoy) == 8 and man[:1] in ("M", "C"):
                out.append((f, nombre, f"Tarde de 8h ({hoy}) seguida de {man}"))
    return sorted(out)


def revisar_cuadro(datos):
    """Problemas del cuadro en la semana: [(fecha o None, médico, problema)].

    - turnos que se cruzan en el tiempo (densitometría no cuenta: se lee en paralelo);
    - nombres incompletos ("Benitez…") o sin el formato "Apellidos, Nombres";
    - códigos que no se reconocen, con médico y día.
    """
    out = []
    tramos = {}
    for k, dia in enumerate(datos["dias"]):
        for srv, a, b, _, n in datos["turnos"][k + 1]:
            tramos.setdefault(n, []).append((k * 24 + a, k * 24 + b, srv, dia))
    for n, ts in tramos.items():
        ts.sort()
        for (a1, b1, s1, d1), (a2, b2, s2, d2) in zip(ts, ts[1:]):
            if a2 < b1 and s1 != s2:
                out.append((d2, n, f"Se cruzan {SERVICIOS.get(s1, s1)} ({hora(a1 % 24)} - {hora(b1 % 24)}) "
                               f"y {SERVICIOS.get(s2, s2)} ({hora(a2 % 24)} - {hora(b2 % 24)})"))
    en_semana = set(tramos) | {n for d in datos["disp"] for n in d.values()} | {n for d in datos["densito"] for n in d}
    for n in sorted(en_semana):
        if n.endswith(("…", "...")) or "," not in n:
            out.append((None, n, "Nombre incompleto en el cuadro (se publica tal cual)"))
    for cod, donde in sorted(datos["desconocidos_detalle"].items()):
        for n, f in donde:
            out.append((f, n, f"Código «{cod}» no reconocido: no aparece en la planilla"))
    return sorted(out, key=lambda x: (x[0] or datetime.min.date(), x[1]))


def equidad(ws, lunes=None):
    """Por médico: horas de la semana (si se da `lunes`) + totales del mes que ya calcula el cuadro."""
    fechas, horas = fechas_mensual(ws), horas_codigo(ws)
    cab = {str(ws.cell(2, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
    meses = {f.month for f in fechas}
    mes = max(meses, key=lambda m: sum(f.month == m for f in fechas))  # el mes del cuadro
    filas_disp = {}
    for r in range(4, ws.max_row + 1):
        n = ws.cell(r, 5).value
        if isinstance(n, str) and str(ws.cell(r, 4).value or "").strip().lower() == "disp":
            filas_disp.setdefault(n.strip(), []).append(r)

    def num(r, col):
        v = ws.cell(r, cab[col]).value if col in cab else None
        return v if isinstance(v, (int, float)) else None

    out = []
    for nombre, r in filas_medicos(ws).items():
        semana = lunes and sum(horas.get(str(ws.cell(r, fechas[lunes + timedelta(k)]).value or "").strip(), 0)
                               for k in range(7) if lunes + timedelta(k) in fechas)
        disp = sum(1 for rr in filas_disp.get(nombre, []) for f, c in fechas.items()
                   if f.month == mes and str(ws.cell(rr, c).value or "").strip().lower() == "disp")
        out.append(dict(medico=nombre, semana=semana, jornada=num(r, "JORNADA"), mes=num(r, "TOTAL"),
                        meta=num(r, "HORAS A"), pendiente=num(r, "PNTE"), noches=num(r, "NC"),
                        findes_libres=num(r, "FDS"), domingos=num(r, "D"), festivos=num(r, "Fes"),
                        disponibilidades=disp))
    return out


# ---------------------------------------------------------------- agendas
def agenda_por_medico(datos):
    """{médico: [(inicio datetime, fin datetime, descripción)]} de la semana generada."""
    eventos = {}
    for k, dia in enumerate(datos["dias"]):
        i = k + 1
        base = datetime.combine(dia, datetime.min.time())
        tramos = {}
        for srv, a, b, rem, n in datos["turnos"][i]:
            tramos.setdefault((n, srv, rem), []).append((a, b))
        for (n, srv, rem), ts in tramos.items():
            for a, b in unir(sorted(ts)):
                desc = SERVICIOS.get(srv, srv) + (" (remoto)" if rem else "")
                eventos.setdefault(n, []).append((base + timedelta(hours=a), base + timedelta(hours=b), desc))
        for n in datos["densito"][i]:
            eventos.setdefault(n, []).append((base + timedelta(hours=7), base + timedelta(hours=19),
                                              SERVICIOS["DENSITO"]))
        for srv, n in datos["disp"][i].items():
            eventos.setdefault(n, []).append((base, None, f"Disponibilidad {SERVICIOS.get(srv, srv)}"))
    return {n: sorted(ev, key=lambda e: (e[0], e[1] is not None)) for n, ev in sorted(eventos.items())}


def unir(tramos):
    out = []
    for a, b in tramos:
        if out and out[-1][1] == a:
            out[-1] = (out[-1][0], b)
        else:
            out.append((a, b))
    return out


def texto_agenda(nombre, eventos, lunes):
    lineas = [f"Turnos de {nombre} · semana del {lunes:%d/%m/%Y}"]
    for ini, fin, desc in eventos:
        if fin is None:
            cuando = "todo el día"
        else:
            cuando = f"{hora(ini.hour)} - {hora(fin.hour)}" + (" (día siguiente)" if fin.date() > ini.date() else "")
        lineas.append(f"• {fecha_corta(ini.date())}: {cuando} — {desc}")
    return "\n".join(lineas)


def ics(nombre, eventos):
    """Calendario iCalendar (hora local) importable en Outlook."""
    def esc(t):
        return t.replace("\\", "\\\\").replace(",", "\\,").replace(";", "\\;")
    ahora = datetime.now().strftime("%Y%m%dT%H%M%S")
    lineas = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Planilla Semanal//ES", "CALSCALE:GREGORIAN"]
    uid = re.sub(r"\W", "", nombre)
    for k, (ini, fin, desc) in enumerate(eventos):
        lineas += ["BEGIN:VEVENT", f"UID:{ini:%Y%m%dT%H%M}-{k}-{uid}@planilla",
                   f"DTSTAMP:{ahora}", f"SUMMARY:{esc(desc)}"]
        if fin is None:
            lineas += [f"DTSTART;VALUE=DATE:{ini:%Y%m%d}", f"DTEND;VALUE=DATE:{ini + timedelta(1):%Y%m%d}"]
        else:
            lineas += [f"DTSTART:{ini:%Y%m%dT%H%M%S}", f"DTEND:{fin:%Y%m%dT%H%M%S}"]
        lineas.append("END:VEVENT")
    lineas.append("END:VCALENDAR")
    return "\r\n".join(lineas) + "\r\n"


def exportar_agendas(datos, carpeta, lunes):
    """Escribe un .ics y un .txt por médico; devuelve la carpeta creada."""
    destino = Path(carpeta) / f"Agendas {nombre_semana(lunes)}"
    destino.mkdir(exist_ok=True)
    for nombre, ev in agenda_por_medico(datos).items():
        archivo = re.sub(r'[\\/:*?"<>|]', "", nombre).strip()
        (destino / f"{archivo}.ics").write_text(ics(nombre, ev), encoding="utf-8")
        (destino / f"{archivo}.txt").write_text(texto_agenda(nombre, ev, lunes), encoding="utf-8")
    return destino


# ---------------------------------------------------------------- cambios entre versiones
def asignaciones(wb):
    """{clave legible: (nombre, celda)} de todas las celdas de nombres de una planilla semanal."""
    out = {}
    eco = wb["Ecografía - ECO"]
    puesto = ""
    for r in range(3, 35):
        # C y D están combinadas: el valor solo está en la primera fila del bloque
        puesto = eco.cell(r, 4).value or eco.cell(r, 3).value or puesto
        etiqueta = f"{puesto} {eco.cell(r, 5).value or ''}".replace("\n", " ").strip()
        for k in range(7):
            c = eco.cell(r, 6 + k)
            out[(k, "Ecografía", etiqueta)] = (c.value, c.coordinate, "Ecografía - ECO")
    for hoja in HOJAS_LECTURA:
        ws = wb[hoja]
        for col in (3, 7):
            titulo = ws.cell(3, col).value
            if not titulo:
                continue
            dia, vistos = -1, {}
            for r in range(4, ws.max_row + 1):
                v = ws.cell(r, col).value
                if hasattr(v, "year") or (isinstance(v, str) and v.startswith("=+")):
                    dia += 1
                    continue
                if dia < 0 or (v is None and ws.cell(r, col + 1).value is None):
                    continue
                base = (dia, titulo, f"{v} {ws.cell(r, col + 2).value or ''}".strip())
                n = vistos[base] = vistos.get(base, 0) + 1
                c = ws.cell(r, col + 1)
                out[base + ((n,) if n > 1 else ())] = (c.value, c.coordinate, hoja)
    return out


def version_anterior(carpeta, lunes):
    """La última planilla guardada de esa semana (SEMANA ....xlsx, v2, v3...) o None."""
    base = nombre_semana(lunes)
    previas = [p for p in Path(carpeta).glob(f"{base}*.xlsx")
               if re.fullmatch(rf"{re.escape(base)}( v\d+)?", p.stem)]
    return max(previas, key=lambda p: int(p.stem[len(base) + 2:] or 1), default=None)


def comparar_libros(viejo, nuevo, lunes):
    """Marca en amarillo las celdas cambiadas de `nuevo` y devuelve [texto del cambio]."""
    a, b = asignaciones(viejo), asignaciones(nuevo)
    cambios = []
    for clave in sorted(set(a) | set(b), key=lambda k: (k[0], str(k[1:]))):
        antes = (a.get(clave) or (None,))[0]
        despues, celda, hoja = b.get(clave) or (None, None, None)
        if (antes or None) == (despues or None):
            continue
        dia, area, puesto = clave[:3]
        cambios.append(f"{fecha_corta(lunes + timedelta(dia))} · {area} · {puesto}: "
                       f"{antes or '(vacío)'} → {despues or '(vacío)'}")
        if celda:
            nuevo[hoja][celda].fill = AMARILLO
    if cambios:
        ws = nuevo.create_sheet("Cambios", 0)
        ws["A1"] = "Cambios respecto a la versión anterior"
        ws["A1"].font = Font(bold=True, size=13)
        for k, t in enumerate(cambios, 3):
            ws.cell(k, 1, t)
        ws.column_dimensions["A"].width = 110
    return cambios


METRICAS_MES = ["mes", "meta", "pendiente", "noches", "domingos", "festivos", "findes_libres", "disponibilidades"]


def equidad_meses(hojas):
    """Suma por médico las métricas mensuales de varios cuadros.

    hojas: [(etiqueta del mes, hoja "Cambios ...")]. Devuelve [{medico, meses, <métricas>}].
    """
    acum = {}
    for etiqueta, ws in hojas:
        for e in equidad(ws):
            fila = acum.setdefault(e["medico"], {"medico": e["medico"], "meses": [], **dict.fromkeys(METRICAS_MES)})
            fila["meses"].append(etiqueta)
            for k in METRICAS_MES:
                if e[k] is not None:
                    fila[k] = (fila[k] or 0) + e[k]
    return sorted(acum.values(), key=lambda f: f["medico"])


# ---------------------------------------------------------------- versión imprimible


def texto_celda(ws, celda):
    """Texto que ve el usuario (las fechas y la fila de fórmulas de ECO se traducen)."""
    v = celda.value
    if isinstance(v, datetime):
        return f"{DIAS_LARGOS[v.weekday()]} {v.day} {MESES_CORTOS[v.month - 1]}"
    if isinstance(v, str) and v.startswith("="):
        # Solo la fila 1 de ECO tiene fórmulas visibles: día de la semana (+ FESTIVO)
        abajo = ws.cell(celda.row + 1, celda.column).value
        if isinstance(abajo, datetime):
            return DIAS_LARGOS[abajo.weekday()].upper() + (" · FESTIVO" if "FESTIVO" in v else "")
        return ""
    return "" if v is None else str(v)


def paleta_tema(wb):
    """Colores del tema del libro, en el orden de los índices de Excel (0 = lt1, 1 = dk1, 2 = lt2...)."""
    xml = wb.loaded_theme.decode() if isinstance(wb.loaded_theme, bytes) else (wb.loaded_theme or "")
    def uno(tag):
        m = re.search(rf"<a:{tag}>.*?(?:lastClr|val)=\"([0-9A-Fa-f]{{6}})\"", xml, re.S)
        return m.group(1).upper() if m else "000000"
    return [uno(t) for t in ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
                              "accent5", "accent6")]


def color_hex(color, tema=None):
    """'#RRGGBB' de un color de openpyxl (rgb o de tema con su tinte); None si no aplica."""
    if color is None:
        return None
    if color.type == "rgb":
        rgb = color.rgb
        if isinstance(rgb, str) and len(rgb) == 8 and rgb not in ("00000000", "FF000000"):
            return "#" + rgb[2:]
        return None
    if color.type == "theme" and tema and color.theme < len(tema):
        canales = [int(tema[color.theme][i:i + 2], 16) for i in (0, 2, 4)]
        t = color.tint or 0
        canales = [round(c + (255 - c) * t) if t > 0 else round(c * (1 + t)) for c in canales]
        return "#" + "".join(f"{c:02X}" for c in canales)
    return None


def html_planilla(wb, lunes, hojas):
    """Página HTML de la planilla para imprimir o guardar como PDF desde el navegador.

    hojas: [(nombre de hoja, título, primera columna a mostrar)].
    """
    from html import escape
    partes = [f"""<!doctype html><html lang="es"><head><meta charset="utf-8">
<title>{nombre_semana(lunes)}</title><style>
body {{ font: 11px "Segoe UI", system-ui, sans-serif; color: #1F2933; margin: 24px; }}
h1 {{ font-size: 18px; margin: 0 0 4px; }} h2 {{ font-size: 15px; margin: 24px 0 8px; }}
p.nota {{ color: #6B7280; margin: 0 0 12px; }}
table {{ border-collapse: collapse; margin-bottom: 8px; }}
td {{ border: 1px solid #D5DAE0; padding: 3px 6px; vertical-align: middle; }}
td.vacia {{ border: none; }}
section {{ page-break-after: always; }} section:last-child {{ page-break-after: auto; }}
@media print {{ body {{ margin: 0; }} p.nota {{ display: none; }} @page {{ size: landscape; margin: 10mm; }} }}
</style></head><body>
<h1>{nombre_semana(lunes)}</h1>
<p class="nota">Para imprimir o guardar como PDF: Ctrl+P (en Mac, Cmd+P).</p>"""]
    tema = paleta_tema(wb)
    for hoja, titulo, col_ini in hojas:
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        tapadas, spans = set(), {}
        for rng in ws.merged_cells.ranges:
            if rng.min_col >= col_ini:
                spans[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
                tapadas |= {(r, c) for r in range(rng.min_row, rng.max_row + 1)
                            for c in range(rng.min_col, rng.max_col + 1)} - {(rng.min_row, rng.min_col)}
        filas = []
        for row in ws.iter_rows(min_col=col_ini):
            celdas = []
            for cel in row:
                if (cel.row, cel.column) in tapadas:
                    continue
                texto = texto_celda(ws, cel)
                fondo = color_hex(cel.fill.fgColor, tema) if cel.fill and cel.fill.fill_type == "solid" else None
                letra = color_hex(cel.font.color, tema) if cel.font else None
                if letra == "#000000":
                    letra = None
                estilo = []
                if fondo:
                    estilo.append(f"background:{fondo}")
                if cel.font and cel.font.b:
                    estilo.append("font-weight:600")
                if letra:
                    estilo.append(f"color:{letra}")
                rs, cs = spans.get((cel.row, cel.column), (1, 1))
                attrs = (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
                if estilo:
                    attrs += f' style="{";".join(estilo)}"'
                if not (texto or fondo or (rs, cs) != (1, 1)):
                    attrs += ' class="vacia"'
                celdas.append(f"<td{attrs}>{escape(texto).replace(chr(10), '<br>')}</td>")
            if any('class="vacia"' not in c for c in celdas):
                filas.append("<tr>" + "".join(celdas) + "</tr>")
        partes.append(f"<section><h2>{escape(titulo)}</h2><table>{''.join(filas)}</table></section>")
    partes.append("</body></html>")
    return "\n".join(partes)


# ---------------------------------------------------------------- resumen y semanas
def semana_texto(lunes):
    dom = lunes + timedelta(6)
    ini = f"{lunes.day}" if lunes.month == dom.month else f"{lunes.day} {MESES_CORTOS[lunes.month - 1]}"
    return f"{ini} – {dom.day} {MESES_CORTOS[dom.month - 1]} {dom.year}"


def semana_por_defecto(lunes, hoy):
    """La primera semana que no ha pasado; si todas pasaron, la última."""
    esta = hoy - timedelta(days=hoy.weekday())
    return next((l for l in lunes if l >= esta), lunes[-1])


def resumen_semana(lunes, datos, resumen, alertas, revision, n_agendas, version_anterior, cambios):
    """Lo que muestra la pestaña Resumen: tarjetas, cobertura por servicio y día, y el aviso «Revisa antes de publicar»."""
    dom = lunes + timedelta(6)
    pend = resumen["pendientes"]
    cruces = [r for r in revision if r[2].startswith("Se cruzan")]
    servicios = sorted({SERVICIOS.get(s, s) for _, _, s, _ in pend})
    fest = resumen["festivos"]
    if version_anterior is None:
        camb = ("—", "info", "Aún no hay versión guardada")
    else:
        camb = (len(cambios), "warn" if cambios else "ok", f"vs. «{version_anterior}»")
    tarjetas = [
        ("Franjas sin médico", len(pend), "warn" if pend else "ok", ", ".join(servicios[:4]) or "Todo cubierto"),
        ("Alertas de descanso", len(alertas), "warn" if alertas else "ok",
         f"{len({a[1] for a in alertas})} médico(s)" if alertas else "Sin alertas"),
        ("Problemas en el cuadro", len(revision), "error" if cruces else ("warn" if revision else "ok"),
         f"{len(cruces)} cruce(s) de turnos" if cruces else ("Revisar" if revision else "Sin problemas")),
        ("Cambios desde la última versión", *camb),
        ("Médicos con turnos", n_agendas, "info", "Ver agendas"),
        ("Festivos", len(fest), "info", ", ".join(fecha_corta(f) for f in fest) or "Ninguno"),
    ]
    cobertura = []
    for srv in COBERTURA:
        dias = []
        for k, dia in enumerate(datos["dias"]):
            i = k + 1
            n = len(datos["densito"][i]) if srv == "DENSITO" else len({t[4] for t in datos["turnos"][i] if t[0] == srv})
            dias.append(dict(n=n, faltan=[fr for f, fr, s, _ in pend if s == srv and f == dia]))
        cobertura.append(dict(servicio=SERVICIOS.get(srv, srv), dias=dias))
    partes = []
    if cruces:
        partes.append(f"{len(cruces)} turno(s) que se cruzan en el cuadro")
    if resumen["desconocidos"]:
        partes.append("códigos no reconocidos (no se publican): " + ", ".join(resumen["desconocidos"]))
    return dict(
        semana=f"Semana del {lunes.day} de {MESES_LARGOS[lunes.month - 1]} al {dom.day} de {MESES_LARGOS[dom.month - 1]}",
        tarjetas=[dict(titulo=t, valor=v, nivel=n, detalle=d) for t, v, n, d in tarjetas],
        cobertura=cobertura, dias=DIAS,
        aviso=("Revisa antes de publicar: " + "; ".join(partes) + ".") if partes else "")
