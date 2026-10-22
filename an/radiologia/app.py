"""Planilla Semanal: ventana para generar la planilla semanal desde el cuadro mensual."""
import os
import subprocess
import sys
import traceback
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from PySide6.QtCore import QFileSystemWatcher, QRectF, QSettings, QSize, Qt, QTimer
from PySide6.QtGui import (QAction, QColor, QFont, QGuiApplication, QIcon, QKeySequence, QPainter,
                           QPen, QPixmap)
from PySide6.QtWidgets import (QAbstractItemView, QApplication, QComboBox, QFileDialog, QFrame,
                               QGridLayout, QHBoxLayout, QHeaderView, QLabel, QLineEdit,
                               QListWidget, QListWidgetItem, QMainWindow, QMessageBox,
                               QPlainTextEdit, QPushButton, QSplitter, QStackedWidget, QTableWidget,
                               QTableWidgetItem, QTabWidget, QToolButton, QVBoxLayout, QWidget)

from . import analisis
from . import generar_semana as gs

LOG = Path(__file__).resolve().parents[2] / "errores.log"
MESES = analisis.MESES_CORTOS
ACENTO = "#0F6E8C"
PESTANAS_PREVIA = {"Ecografía - ECO": "Ecografía", "Lectura - Neuro": "Neuro", "Lectura - MSK": "MSK",
                   "Lectura - Abdomen": "Abdomen", "Lectura - CARDIO-TORAX AMBULATO": "Tórax y Corazón",
                   "Lectura - Mamografia-Densitomet": "Mamografía y Densito", "Cambios": "Cambios"}
SECCIONES = ["Resumen", "Vista previa", "Alertas", "Equidad y horas", "Agendas", "Cambios"]
INICIO, PREVIA, ALERTAS, EQUIDAD, AGENDAS, CAMBIOS = range(len(SECCIONES))
# Servicios en el orden de la tabla de cobertura
COBERTURA = ["URG", "ECO1", "ECO2", "ECO3", "ECO4", "ECO5", "PM", "CJ1", "CJ2",
             "NEURO", "MSK", "ABD", "TORAX", "CARDIO", "MAMA", "DENSITO"]

ESTILO = f"""
QMainWindow, QWidget#fondo {{ background: #F4F6F8; }}
QLabel {{ color: #1F2933; }}
QLabel#muted, QLabel#sub {{ color: #6B7280; }}
QLabel#marca {{ font-size: 16px; font-weight: 700; }}
QLabel#seccion {{ font-size: 18px; font-weight: 650; }}
QLabel#campo {{ color: #6B7280; font-size: 11px; font-weight: 600; letter-spacing: .4px; }}
QFrame#cabecera {{ background: white; border-bottom: 1px solid #E3E6EA; }}
QListWidget#nav {{ background: transparent; border: none; outline: 0; font-size: 13px; }}
QListWidget#nav::item {{ padding: 9px 12px; border-radius: 7px; color: #3E4C59; margin: 1px 0; }}
QListWidget#nav::item:hover {{ background: #E9EDF1; }}
QListWidget#nav::item:selected {{ background: #DCECF1; color: {ACENTO}; font-weight: 600; }}
QFrame#tarjeta, QFrame#panel {{ background: white; border: 1px solid #E3E6EA; border-radius: 10px; }}
QFrame#tarjeta:hover {{ border-color: #B7C9D3; }}
QFrame#tarjeta[nivel="warn"] {{ border-left: 4px solid #E0A800; }}
QFrame#tarjeta[nivel="error"] {{ border-left: 4px solid #C53030; }}
QFrame#tarjeta[nivel="ok"] {{ border-left: 4px solid #2F855A; }}
QFrame#tarjeta[nivel="info"] {{ border-left: 4px solid {ACENTO}; }}
QLabel#cifra {{ font-size: 26px; font-weight: 700; }}
QLabel#aviso {{ border-radius: 8px; padding: 10px 12px; }}
QLabel#aviso[nivel="info"] {{ background: #E8F1F8; color: #12466B; }}
QLabel#aviso[nivel="warn"] {{ background: #FFF4D6; color: #6B4E00; }}
QLabel#aviso[nivel="error"] {{ background: #FDE8E8; color: #8E1B1B; }}
QLabel#error {{ background: #FDE8E8; color: #8E1B1B; border-radius: 8px; padding: 10px 14px; }}
QLabel#toast {{ background: #1F2933; color: white; border-radius: 8px; padding: 9px 14px; }}
QFrame#soltar {{ background: white; border: 2px dashed #C5CED6; border-radius: 14px; }}
QFrame#soltar[activo="true"] {{ border-color: {ACENTO}; background: #EEF6F9; }}
QPushButton, QToolButton {{ background: white; border: 1px solid #CBD2D9; border-radius: 6px; padding: 7px 14px; color: #1F2933; }}
QPushButton:hover, QToolButton:hover {{ background: #F0F3F6; }}
QPushButton:disabled, QToolButton:disabled {{ color: #A0A7B0; background: #F5F6F8; }}
QPushButton#primario {{ background: {ACENTO}; border: 1px solid {ACENTO}; color: white; font-weight: 600; padding: 8px 18px; }}
QPushButton#primario:hover {{ background: #0C5C75; }}
QPushButton#primario:disabled {{ background: #9DBFCB; border-color: #9DBFCB; }}
QPushButton#archivo {{ text-align: left; padding: 7px 12px; max-width: 360px; }}
QToolButton#flecha {{ padding: 6px 9px; font-weight: 700; }}
QComboBox, QLineEdit {{ background: white; border: 1px solid #CBD2D9; border-radius: 6px; padding: 6px 10px; }}
QComboBox {{ min-width: 190px; }}
QComboBox::drop-down {{ border: none; width: 22px; }}
QLineEdit:focus, QComboBox:focus {{ border-color: {ACENTO}; }}
QTabWidget::pane {{ border: none; background: white; }}
QTabBar::tab {{ padding: 8px 14px; color: #52606D; border: none; }}
QTabBar::tab:selected {{ color: {ACENTO}; border-bottom: 2px solid {ACENTO}; font-weight: 600; }}
QTableWidget {{ border: none; gridline-color: #EDF0F3; selection-background-color: #D6E9F0; selection-color: #1F2933; }}
QHeaderView::section {{ background: #F7F8FA; border: none; border-bottom: 1px solid #E3E6EA; padding: 6px; color: #52606D; font-weight: 600; }}
QListWidget#medicos, QPlainTextEdit {{ border: 1px solid #E3E6EA; border-radius: 6px; background: white; }}
QListWidget#medicos::item {{ padding: 6px 8px; }}
QListWidget#medicos::item:selected {{ background: #DCECF1; color: #1F2933; }}
QSplitter::handle {{ background: transparent; }}
"""


# ---------------------------------------------------------------- utilidades
def semana_texto(lunes):
    dom = lunes + timedelta(6)
    ini = f"{lunes.day}" if lunes.month == dom.month else f"{lunes.day} {MESES[lunes.month - 1]}"
    return f"{ini} – {dom.day} {MESES[dom.month - 1]} {dom.year}"


def abrir(ruta):
    if sys.platform == "win32":
        os.startfile(ruta)  # solo existe en Windows
    else:
        subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", str(ruta)])


def icono_app():
    """Ícono dibujado (sin archivos): calendario blanco sobre el color de la app."""
    pm = QPixmap(128, 128)
    pm.fill(Qt.transparent)
    p = QPainter(pm)
    p.setRenderHint(QPainter.Antialiasing)
    p.setBrush(QColor(ACENTO))
    p.setPen(Qt.NoPen)
    p.drawRoundedRect(QRectF(4, 4, 120, 120), 26, 26)
    p.setBrush(QColor("white"))
    p.drawRoundedRect(QRectF(26, 34, 76, 66), 9, 9)
    p.setBrush(QColor(ACENTO))
    p.drawRect(QRectF(26, 34, 76, 15))
    p.setPen(QPen(QColor("white"), 6, Qt.SolidLine, Qt.RoundCap))
    p.drawLine(44, 24, 44, 40)
    p.drawLine(84, 24, 84, 40)
    p.setPen(Qt.NoPen)
    for i in range(3):
        for j in range(2):
            p.setBrush(QColor(ACENTO if (i, j) != (1, 1) else "#E0A800"))
            p.drawRoundedRect(QRectF(36 + i * 20, 58 + j * 18, 13, 11), 3, 3)
    p.end()
    return QIcon(pm)


def tabla(columnas, ayudas=None):
    t = QTableWidget(0, len(columnas))
    t.setHorizontalHeaderLabels(columnas)
    t.setEditTriggers(QAbstractItemView.NoEditTriggers)
    t.setSelectionBehavior(QAbstractItemView.SelectRows)
    t.verticalHeader().setVisible(False)
    t.horizontalHeader().setStretchLastSection(True)
    t.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
    for c, ayuda in (ayudas or {}).items():
        t.horizontalHeaderItem(c).setToolTip(ayuda)
    return t


def llenar(t, filas, colores=None, orden=None):
    t.setSortingEnabled(False)
    t.clearContents()
    t.setRowCount(len(filas))
    for r, fila in enumerate(filas):
        for c, v in enumerate(fila):
            it = QTableWidgetItem()
            it.setData(Qt.DisplayRole, "—" if v is None else v)
            if colores and (color := colores(r, c)):
                it.setBackground(QColor(color))
            t.setItem(r, c, it)
    t.resizeColumnsToContents()
    if orden is not None:
        t.setSortingEnabled(True)
        t.sortByColumn(orden, Qt.AscendingOrder)


def color_qt(color, tema):
    h = analisis.color_hex(color, tema)
    return QColor(h) if h else None


def tabla_hoja(ws, col_ini, tema):
    """Copia visual (valores, combinadas, rellenos, negrita) de una hoja de la planilla."""
    cols = max(ws.max_column, col_ini) - col_ini + 1
    t = QTableWidget(ws.max_row, cols)
    t.setEditTriggers(QAbstractItemView.NoEditTriggers)
    t.horizontalHeader().setVisible(False)
    t.verticalHeader().setVisible(False)
    t.setShowGrid(False)
    negrita = QFont()
    negrita.setBold(True)
    for row in ws.iter_rows(min_col=col_ini):
        for cel in row:
            solido = cel.fill is not None and cel.fill.fill_type == "solid"
            if cel.value is None and not solido:
                continue
            it = QTableWidgetItem(analisis.texto_celda(ws, cel))
            if solido and (bg := color_qt(cel.fill.fgColor, tema)):
                it.setBackground(bg)
            if cel.font is not None:
                if cel.font.b:
                    it.setFont(negrita)
                if (fg := color_qt(cel.font.color, tema)) is not None:
                    it.setForeground(fg)
            t.setItem(cel.row - 1, cel.column - col_ini, it)
    for rng in ws.merged_cells.ranges:
        if rng.min_col >= col_ini:
            t.setSpan(rng.min_row - 1, rng.min_col - col_ini, rng.max_row - rng.min_row + 1,
                      rng.max_col - rng.min_col + 1)
    for c in range(cols):
        t.setColumnWidth(c, 190 if c else 230)
    t.resizeRowsToContents()
    return t


class ocupado:
    def __enter__(self):
        QApplication.setOverrideCursor(Qt.WaitCursor)

    def __exit__(self, *exc):
        QApplication.restoreOverrideCursor()


class Tarjeta(QFrame):
    """Cifra grande + etiqueta; al hacer clic lleva a una sección."""

    def __init__(self, titulo, al_clic):
        super().__init__(objectName="tarjeta")
        self.setCursor(Qt.PointingHandCursor)
        self.al_clic = al_clic
        v = QVBoxLayout(self)
        v.setContentsMargins(16, 12, 16, 12)
        v.setSpacing(2)
        self.cifra = QLabel("—", objectName="cifra")
        self.titulo = QLabel(titulo, objectName="muted")
        self.detalle = QLabel("", objectName="muted")
        self.detalle.setWordWrap(True)
        for w in (self.cifra, self.titulo, self.detalle):
            v.addWidget(w)

    def poner(self, cifra, nivel, detalle=""):
        self.cifra.setText(str(cifra))
        self.detalle.setText(detalle)
        self.setProperty("nivel", nivel)
        self.style().unpolish(self)
        self.style().polish(self)

    def mousePressEvent(self, e):
        self.al_clic()


# ---------------------------------------------------------------- ventana
class Ventana(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Planilla Semanal · Radiología")
        self.setWindowIcon(icono_app())
        self.setAcceptDrops(True)
        self.cfg = QSettings("Radiologia", "PlanillaSemanal")
        self.ws = self.ruta = self.wb = self.datos = self.lunes = None
        self.agendas = {}
        self._previa_pendiente = False
        self._cache_previa = {}  # ruta -> (mtime, workbook) de la versión anterior guardada

        self.vigia = QFileSystemWatcher(self)
        self.vigia.fileChanged.connect(lambda _: self._recarga.start())
        self._recarga = QTimer(self, singleShot=True, interval=800, timeout=self.recargar)

        fondo = QWidget(objectName="fondo")
        self.setCentralWidget(fondo)
        raiz = QVBoxLayout(fondo)
        raiz.setContentsMargins(0, 0, 0, 0)
        raiz.setSpacing(0)
        raiz.addWidget(self._cabecera())

        cuerpo = QHBoxLayout()
        cuerpo.setContentsMargins(16, 16, 20, 16)
        cuerpo.setSpacing(16)
        self.nav = QListWidget(objectName="nav")
        self.nav.setFixedWidth(190)
        self.nav.addItems(SECCIONES)
        self.nav.currentRowChanged.connect(self._ir)
        cuerpo.addWidget(self.nav)

        der = QVBoxLayout()
        der.setSpacing(12)
        self.aviso = QLabel(objectName="aviso", wordWrap=True)
        self.aviso.setTextInteractionFlags(Qt.TextSelectableByMouse)
        self.aviso.hide()
        der.addWidget(self.aviso)
        self.paginas = QStackedWidget()
        for crear in (self._pag_resumen, self._pag_previa, self._pag_alertas, self._pag_equidad,
                      self._pag_agendas, self._pag_cambios):
            self.paginas.addWidget(crear())
        der.addWidget(self.paginas, 1)
        cuerpo.addLayout(der, 1)

        self.principal = QWidget()
        self.principal.setLayout(cuerpo)
        self.pila = QStackedWidget()
        self.pila.addWidget(self._pag_vacia())
        self.pila.addWidget(self.principal)
        raiz.addWidget(self.pila, 1)

        self.toast = QLabel(self, objectName="toast")
        self.toast.hide()
        self._toast_timer = QTimer(self, singleShot=True, interval=4500, timeout=self.toast.hide)

        self._atajos()
        self._habilitar(False)
        if geo := self.cfg.value("geometria"):
            self.restoreGeometry(geo)
        else:
            self.resize(1280, 840)
        ultimo = self.cfg.value("ultimo_archivo", "")
        if ultimo and Path(ultimo).exists():
            self.cargar(ultimo)

    # ---------------------------------------------------------------- estructura
    def _cabecera(self):
        cab = QFrame(objectName="cabecera")
        h = QHBoxLayout(cab)
        h.setContentsMargins(20, 12, 20, 12)
        h.setSpacing(10)
        logo = QLabel()
        logo.setPixmap(icono_app().pixmap(QSize(34, 34)))
        h.addWidget(logo)
        marca = QVBoxLayout()
        marca.setSpacing(0)
        marca.addWidget(QLabel("Planilla Semanal", objectName="marca"))
        marca.addWidget(QLabel("Radiología", objectName="sub"))
        h.addLayout(marca)
        h.addSpacing(24)

        self.btn_archivo = QPushButton("Elegir cuadro mensual…", objectName="archivo")
        self.btn_archivo.setToolTip("Cuadro mensual de turnos (Ctrl+O). También puedes arrastrarlo a la ventana.")
        self.btn_archivo.clicked.connect(self.elegir_archivo)
        h.addWidget(self.btn_archivo)

        self.btn_ant = QToolButton(objectName="flecha", text="‹", toolTip="Semana anterior (Alt+←)")
        self.btn_sig = QToolButton(objectName="flecha", text="›", toolTip="Semana siguiente (Alt+→)")
        self.combo = QComboBox()
        self.combo.currentIndexChanged.connect(self.cambiar_semana)
        self.btn_ant.clicked.connect(lambda: self._mover_semana(-1))
        self.btn_sig.clicked.connect(lambda: self._mover_semana(1))
        for w in (self.btn_ant, self.combo, self.btn_sig):
            h.addWidget(w)
        h.addStretch(1)

        self.btn_imprimir = QPushButton("Imprimir", toolTip="Versión para imprimir o guardar como PDF (Ctrl+P)")
        self.btn_imprimir.clicked.connect(self.imprimir)
        self.btn_agendas = QPushButton("Exportar agendas", toolTip="Un .ics (Outlook) y un .txt por médico")
        self.btn_agendas.clicked.connect(self.exportar_agendas)
        self.btn_carpeta = QPushButton("Abrir carpeta", toolTip="Carpeta donde se guardan las planillas")
        self.btn_carpeta.clicked.connect(lambda: abrir(Path(self.ruta).parent))
        self.btn_generar = QPushButton("Generar Excel", objectName="primario",
                                       toolTip="Guarda la planilla junto al cuadro y la abre (Ctrl+G)")
        self.btn_generar.clicked.connect(self.generar)
        for b in (self.btn_carpeta, self.btn_imprimir, self.btn_agendas, self.btn_generar):
            h.addWidget(b)
        return cab

    def _pag_vacia(self):
        caja = QWidget()
        v = QVBoxLayout(caja)
        v.setContentsMargins(60, 40, 60, 60)
        self.soltar = QFrame(objectName="soltar")
        s = QVBoxLayout(self.soltar)
        s.setAlignment(Qt.AlignCenter)
        s.setSpacing(10)
        ic = QLabel(alignment=Qt.AlignCenter)
        ic.setPixmap(icono_app().pixmap(QSize(72, 72)))
        titulo = QLabel("Arrastra aquí el cuadro mensual", objectName="seccion", alignment=Qt.AlignCenter)
        sub = QLabel("Archivo de Excel con la hoja «Cambios …» y fechas completas en la fila 3.",
                     objectName="muted", alignment=Qt.AlignCenter)
        btn = QPushButton("Elegir archivo…", objectName="primario")
        btn.clicked.connect(self.elegir_archivo)
        self.error_carga = QLabel(objectName="error", wordWrap=True, alignment=Qt.AlignCenter)
        self.error_carga.setMaximumWidth(620)
        self.error_carga.hide()
        for w in (ic, titulo, sub):
            s.addWidget(w)
        s.addWidget(self.error_carga, 0, Qt.AlignCenter)
        s.addWidget(btn, 0, Qt.AlignCenter)
        v.addWidget(self.soltar, 1)
        return caja

    def _titulo(self, texto, ayuda):
        caja = QVBoxLayout()
        caja.setSpacing(2)
        caja.addWidget(QLabel(texto, objectName="seccion"))
        caja.addWidget(QLabel(ayuda, objectName="muted", wordWrap=True))
        return caja

    def _pag_resumen(self):
        pag = QWidget()
        v = QVBoxLayout(pag)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(14)
        self.lbl_resumen = QLabel(objectName="seccion")
        self.lbl_destino = QLabel(objectName="muted")
        cab = QVBoxLayout()
        cab.setSpacing(2)
        cab.addWidget(self.lbl_resumen)
        cab.addWidget(self.lbl_destino)
        v.addLayout(cab)

        grid = QGridLayout()
        grid.setSpacing(12)
        self.t_pend = Tarjeta("Franjas sin médico", lambda: self._ir_a(ALERTAS))
        self.t_desc = Tarjeta("Alertas de descanso", lambda: self._ir_a(ALERTAS))
        self.t_rev = Tarjeta("Problemas en el cuadro", lambda: self._ir_a(ALERTAS))
        self.t_camb = Tarjeta("Cambios desde la última versión", lambda: self._ir_a(CAMBIOS))
        self.t_med = Tarjeta("Médicos con turnos", lambda: self._ir_a(AGENDAS))
        self.t_fest = Tarjeta("Festivos", lambda: self._ir_a(PREVIA))
        for i, t in enumerate((self.t_pend, self.t_desc, self.t_rev, self.t_camb, self.t_med, self.t_fest)):
            grid.addWidget(t, i // 3, i % 3)
        v.addLayout(grid)

        panel = QFrame(objectName="panel")
        pv = QVBoxLayout(panel)
        pv.setContentsMargins(14, 12, 14, 12)
        pv.addWidget(QLabel("COBERTURA DE LA SEMANA", objectName="campo"))
        pv.addWidget(QLabel("Médicos asignados por servicio y día. En amarillo, los días con franjas sin médico "
                            "(pasa el mouse para ver cuáles).", objectName="muted", wordWrap=True))
        self.t_cob = tabla(["Servicio"] + analisis.DIAS)
        self.t_cob.horizontalHeader().setStretchLastSection(False)
        self.t_cob.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.t_cob.setSelectionMode(QAbstractItemView.NoSelection)
        pv.addWidget(self.t_cob, 1)
        v.addWidget(panel, 1)
        return pag

    def _pag_previa(self):
        pag = QFrame(objectName="panel")
        v = QVBoxLayout(pag)
        v.setContentsMargins(6, 6, 6, 6)
        self.previa = QTabWidget()
        self.previa.setDocumentMode(True)
        v.addWidget(self.previa)
        return pag

    def _pag_alertas(self):
        pag = QWidget()
        v = QVBoxLayout(pag)
        v.setContentsMargins(0, 0, 0, 0)
        v.addLayout(self._titulo("Alertas", "Revisa esto antes de publicar. Los problemas se corrigen en el "
                                 "cuadro mensual de Excel; al guardarlo, la app se actualiza sola."))
        self.t_revision = tabla(["Fecha", "Médico", "Problema"])
        self.t_alertas = tabla(["Fecha", "Médico", "Regla"])
        self.t_pendientes = tabla(["Fecha", "Servicio", "Franja", "Qué falta"])
        for titulo, t in (("PROBLEMAS EN EL CUADRO MENSUAL", self.t_revision),
                          ("REGLAS DE DESCANSO", self.t_alertas),
                          ("FRANJAS SIN MÉDICO", self.t_pendientes)):
            panel = QFrame(objectName="panel")
            pv = QVBoxLayout(panel)
            pv.setContentsMargins(12, 10, 12, 10)
            pv.addWidget(QLabel(titulo, objectName="campo"))
            pv.addWidget(t)
            v.addWidget(panel, 1)
        return pag

    def _pag_equidad(self):
        pag = QWidget()
        v = QVBoxLayout(pag)
        v.setContentsMargins(0, 0, 0, 0)
        v.addLayout(self._titulo("Equidad y horas", "Carga de cada radiólogo. Haz clic en un encabezado para "
                                 "ordenar; pasa el mouse sobre él para ver qué significa."))
        barra = QHBoxLayout()
        self.modo_eq = QComboBox()
        self.modo_eq.addItems(["Este mes", "Varios meses (acumulado)"])
        self.modo_eq.currentIndexChanged.connect(self._mostrar_meses)
        self.btn_meses = QPushButton("Elegir cuadros de otros meses…")
        self.btn_meses.clicked.connect(self.elegir_meses)
        self.lbl_meses = QLabel(objectName="muted")
        for w in (self.modo_eq, self.btn_meses):
            barra.addWidget(w)
        barra.addWidget(self.lbl_meses, 1)
        v.addLayout(barra)
        self.t_equidad = tabla(
            ["Médico", "Horas semana", "Jornada sem.", "Horas mes", "Meta mes", "Pendiente", "Noches URG",
             "Domingos", "Festivos", "Finde libres", "Disponibilidades"],
            {1: "Horas de la semana elegida según la hoja Nom-Hrs (rojo: supera la jornada)",
             5: "Meta del mes menos horas programadas (amarillo: le faltan horas; azul: tiene de más)",
             9: "Fines de semana completos libres en el mes", 10: "Días con disponibilidad en el mes"})
        self.t_meses = tabla(["Médico", "Meses", "Horas", "Meta", "Pendiente", "Noches URG", "Domingos",
                              "Festivos", "Finde libres", "Disponibilidades"])
        self.pila_eq = QStackedWidget()
        for t in (self.t_equidad, self.t_meses):
            panel = QFrame(objectName="panel")
            QVBoxLayout(panel).addWidget(t)
            self.pila_eq.addWidget(panel)
        v.addWidget(self.pila_eq, 1)
        return pag

    def _pag_agendas(self):
        pag = QWidget()
        v = QVBoxLayout(pag)
        v.setContentsMargins(0, 0, 0, 0)
        v.addLayout(self._titulo("Agendas", "Turnos de cada médico para enviar por WhatsApp o correo. "
                                 "«Exportar agendas» crea además un .ics por médico para Outlook."))
        sp = QSplitter()
        izq = QWidget()
        iv = QVBoxLayout(izq)
        iv.setContentsMargins(0, 0, 0, 0)
        self.buscar = QLineEdit(placeholderText="Buscar médico…", clearButtonEnabled=True)
        self.buscar.textChanged.connect(self._filtrar_medicos)
        self.l_medicos = QListWidget(objectName="medicos")
        self.l_medicos.currentTextChanged.connect(self.mostrar_agenda)
        iv.addWidget(self.buscar)
        iv.addWidget(self.l_medicos, 1)
        der = QWidget()
        dv = QVBoxLayout(der)
        dv.setContentsMargins(0, 0, 0, 0)
        self.txt_agenda = QPlainTextEdit(readOnly=True)
        botones = QHBoxLayout()
        botones.addStretch(1)
        b1 = QPushButton("Copiar esta agenda")
        b1.clicked.connect(lambda: self.copiar(self.txt_agenda.toPlainText(), "Agenda copiada."))
        b2 = QPushButton("Copiar todas")
        b2.clicked.connect(self.copiar_todas)
        botones.addWidget(b2)
        botones.addWidget(b1)
        dv.addWidget(self.txt_agenda, 1)
        dv.addLayout(botones)
        sp.addWidget(izq)
        sp.addWidget(der)
        sp.setSizes([300, 900])
        v.addWidget(sp, 1)
        return pag

    def _pag_cambios(self):
        pag = QWidget()
        v = QVBoxLayout(pag)
        v.setContentsMargins(0, 0, 0, 0)
        v.addLayout(self._titulo("Cambios", "Diferencias con la última planilla guardada de esta semana. "
                                 "Al generar, quedan marcadas en amarillo y en la hoja «Cambios»."))
        self.lbl_cambios = QLabel(objectName="muted")
        self.txt_cambios = QPlainTextEdit(readOnly=True)
        btn = QPushButton("Copiar lista de cambios")
        btn.clicked.connect(lambda: self.copiar(self.txt_cambios.toPlainText(), "Cambios copiados."))
        v.addWidget(self.lbl_cambios)
        v.addWidget(self.txt_cambios, 1)
        v.addWidget(btn, 0, Qt.AlignRight)
        return pag

    def _atajos(self):
        for tecla, accion in ((QKeySequence.Open, self.elegir_archivo), ("Ctrl+G", self.generar),
                              (QKeySequence.Print, self.imprimir), ("Alt+Left", lambda: self._mover_semana(-1)),
                              ("Alt+Right", lambda: self._mover_semana(1)), (QKeySequence.Refresh, self.recargar)):
            a = QAction(self, shortcut=QKeySequence(tecla))
            a.triggered.connect(lambda _=False, f=accion: f() if self._puede(f) else None)
            self.addAction(a)
        for i in range(len(SECCIONES)):
            a = QAction(self, shortcut=QKeySequence(f"Ctrl+{i + 1}"))
            a.triggered.connect(lambda _=False, i=i: self._ir_a(i))
            self.addAction(a)

    def _puede(self, f):
        return f == self.elegir_archivo or self.wb is not None

    # ---------------------------------------------------------------- navegación
    def _ir_a(self, i):
        self.nav.setCurrentRow(i)

    def _ir(self, i):
        self.paginas.setCurrentIndex(i)
        if i == PREVIA and self._previa_pendiente:
            self._mostrar_previa()

    def _mover_semana(self, paso):
        i = self.combo.currentIndex() + paso
        if 0 <= i < self.combo.count():
            self.combo.setCurrentIndex(i)

    def _habilitar(self, si):
        for w in (self.btn_imprimir, self.btn_agendas, self.btn_generar, self.btn_carpeta, self.combo,
                  self.btn_ant, self.btn_sig):
            w.setEnabled(si)
        if si:
            i = self.combo.currentIndex()
            self.btn_ant.setEnabled(i > 0)
            self.btn_sig.setEnabled(i < self.combo.count() - 1)

    # ---------------------------------------------------------------- arrastrar y soltar
    def _xlsx(self, evento):
        urls = evento.mimeData().urls() if evento.mimeData().hasUrls() else []
        return next((u.toLocalFile() for u in urls if u.toLocalFile().lower().endswith(".xlsx")), None)

    def _resaltar(self, si):
        self.soltar.setProperty("activo", si)
        self.soltar.style().unpolish(self.soltar)
        self.soltar.style().polish(self.soltar)

    def dragEnterEvent(self, e):
        if self._xlsx(e):
            e.acceptProposedAction()
            self._resaltar(True)

    def dragLeaveEvent(self, e):
        self._resaltar(False)

    def dropEvent(self, e):
        self._resaltar(False)
        if ruta := self._xlsx(e):
            self.cargar(ruta)

    def closeEvent(self, e):
        self.cfg.setValue("geometria", self.saveGeometry())
        super().closeEvent(e)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        self._ubicar_toast()

    # ---------------------------------------------------------------- acciones
    def elegir_archivo(self):
        inicio = str(Path(self.ruta).parent) if self.ruta else str(Path.home())
        ruta, _ = QFileDialog.getOpenFileName(self, "Elegir cuadro mensual", inicio, "Excel (*.xlsx)")
        if ruta:
            self.cargar(ruta)

    def cargar(self, ruta, semana=None):
        if self.vigia.files():
            self.vigia.removePaths(self.vigia.files())
        self.ruta, self.ws, self.wb = ruta, None, None
        self.btn_archivo.setText(Path(ruta).name)
        self.btn_archivo.setToolTip(ruta)
        self.avisar("")
        self._habilitar(False)
        try:
            with ocupado():
                self.ws = gs.cargar_mensual(ruta)
            lunes = gs.lunes_disponibles(self.ws)
            if not lunes:
                raise gs.ErrorPlanilla("El cuadro no tiene ninguna semana completa (lunes a domingo).")
        except gs.ErrorPlanilla as e:
            return self._sin_cuadro(f"«{Path(ruta).name}»: {e}")
        self.vigia.addPath(ruta)
        self.cfg.setValue("ultimo_archivo", ruta)
        self.error_carga.hide()
        hoy = date.today()
        esta = hoy - timedelta(days=hoy.weekday())
        defecto = (lunes.index(semana) if semana in lunes
                   else next((i for i, l in enumerate(lunes) if l >= esta), len(lunes) - 1))
        self.combo.blockSignals(True)
        self.combo.clear()
        for l in lunes:
            self.combo.addItem(semana_texto(l), l.isoformat())  # Qt convertiría date a QDate
        self.combo.setCurrentIndex(defecto)
        self.combo.blockSignals(False)
        self.pila.setCurrentIndex(1)
        if self.nav.currentRow() < 0:
            self._ir_a(INICIO)
        self.cambiar_semana()

    def _sin_cuadro(self, error):
        """Vuelve a la pantalla de inicio mostrando por qué no se pudo cargar."""
        self.ws = self.wb = None
        self.combo.blockSignals(True)
        self.combo.clear()
        self.combo.blockSignals(False)
        self._habilitar(False)
        self.error_carga.setText(error)
        self.error_carga.show()
        self.pila.setCurrentIndex(0)

    def recargar(self):
        """El Excel cambió en disco (o F5): se vuelve a leer manteniendo la semana."""
        if not self.ruta or not Path(self.ruta).exists():
            return
        self.cargar(self.ruta, self.lunes)
        if self.ws is not None:
            self.notificar("Cuadro mensual actualizado.")

    def cambiar_semana(self, *_):
        iso = self.combo.currentData()
        if self.ws is None or not iso:
            return
        self.lunes = date.fromisoformat(iso)
        try:
            with ocupado():
                self.wb, resumen, self.datos = gs.construir(self.ws, self.lunes)
                previa = analisis.version_anterior(Path(self.ruta).parent, self.lunes)
                cambios = analisis.comparar_libros(self._libro_previo(previa), self.wb, self.lunes) if previa else None
                alertas = analisis.alertas(self.ws, self.lunes)
                revision = analisis.revisar_cuadro(self.datos)
                equidad = analisis.equidad(self.ws, self.lunes)
        except gs.ErrorPlanilla as e:
            self._habilitar(False)
            return self.avisar(str(e), "error")
        self._habilitar(True)
        self._previa_pendiente = True
        if self.paginas.currentIndex() == PREVIA:
            self._mostrar_previa()
        self._mostrar_alertas(alertas, resumen["pendientes"], revision)
        self._mostrar_equidad(equidad)
        self._mostrar_agendas()
        self._mostrar_cambios(previa, cambios)
        self._mostrar_meses()
        self._mostrar_resumen(resumen, alertas, revision, previa, cambios)

    def _libro_previo(self, ruta):
        """La versión anterior se relee solo si cambió en disco."""
        mtime = ruta.stat().st_mtime
        guardado = self._cache_previa.get(ruta)
        if not guardado or guardado[0] != mtime:
            guardado = self._cache_previa[ruta] = (mtime, openpyxl.load_workbook(ruta))
        return guardado[1]

    def generar(self):
        destino = gs.ruta_libre(Path(self.ruta).parent, self.lunes)
        try:
            self.wb.save(destino)
        except PermissionError:
            return self.avisar(f"No se pudo guardar «{destino.name}». ¿Está abierto en Excel?", "error")
        abrir(destino)
        self.cambiar_semana()  # la versión guardada pasa a ser la anterior
        self.notificar(f"Guardado: {destino.name}")

    def imprimir(self):
        hojas = [(h, t, 1 if h == "Cambios" else 3) for h, t in PESTANAS_PREVIA.items()]
        destino = Path(self.ruta).parent / f"{gs.nombre_semana(self.lunes)} (para imprimir).html"
        destino.write_text(analisis.html_planilla(self.wb, self.lunes, hojas), encoding="utf-8")
        abrir(destino)
        self.notificar("Se abrió en el navegador: usa Imprimir o Guardar como PDF.")

    def exportar_agendas(self):
        destino = analisis.exportar_agendas(self.datos, Path(self.ruta).parent, self.lunes)
        abrir(destino)
        self.notificar(f"Agendas guardadas en «{destino.name}».")

    def elegir_meses(self):
        inicio = str(Path(self.ruta).parent) if self.ruta else str(Path.home())
        rutas, _ = QFileDialog.getOpenFileNames(self, "Cuadros de otros meses", inicio, "Excel (*.xlsx)")
        if rutas:
            self.cfg.setValue("otros_meses", rutas)
            self.modo_eq.setCurrentIndex(1)
            self._mostrar_meses()

    def copiar(self, texto, msg):
        QGuiApplication.clipboard().setText(texto)
        self.notificar(msg)

    def copiar_todas(self):
        textos = [analisis.texto_agenda(n, ev, self.lunes) for n, ev in self.agendas.items()]
        self.copiar("\n\n".join(textos), f"{len(textos)} agendas copiadas.")

    def avisar(self, texto, nivel="info"):
        if not texto:
            return self.aviso.hide()
        self.aviso.setText(texto)
        self.aviso.setProperty("nivel", nivel)
        self.aviso.style().unpolish(self.aviso)
        self.aviso.style().polish(self.aviso)
        self.aviso.show()

    def notificar(self, texto):
        self.toast.setText(texto)
        self.toast.adjustSize()
        self._ubicar_toast()
        self.toast.show()
        self.toast.raise_()
        self._toast_timer.start()

    def _ubicar_toast(self):
        self.toast.move(self.width() - self.toast.width() - 24, self.height() - self.toast.height() - 20)

    # ---------------------------------------------------------------- vistas
    def _mostrar_resumen(self, resumen, alertas, revision, previa, cambios):
        dom = self.lunes + timedelta(6)
        self.lbl_resumen.setText(f"Semana del {self.lunes.day} de {analisis.MESES_LARGOS[self.lunes.month - 1]} "
                                 f"al {dom.day} de {analisis.MESES_LARGOS[dom.month - 1]}")
        destino = gs.ruta_libre(Path(self.ruta).parent, self.lunes).name
        self.lbl_destino.setText(f"Al generar se guardará como «{destino}» junto al cuadro mensual.")

        pend = resumen["pendientes"]
        cruces = [r for r in revision if r[2].startswith("Se cruzan")]
        servicios = sorted({analisis.SERVICIOS.get(s, s) for _, _, s, _ in pend})
        self.t_pend.poner(len(pend), "warn" if pend else "ok", ", ".join(servicios[:4]) or "Todo cubierto")
        self.t_desc.poner(len(alertas), "warn" if alertas else "ok",
                          f"{len({a[1] for a in alertas})} médico(s)" if alertas else "Sin alertas")
        self.t_rev.poner(len(revision), "error" if cruces else ("warn" if revision else "ok"),
                         f"{len(cruces)} cruce(s) de turnos" if cruces else ("Revisar" if revision else "Sin problemas"))
        if previa is None:
            self.t_camb.poner("—", "info", "Aún no hay versión guardada")
        else:
            self.t_camb.poner(len(cambios), "warn" if cambios else "ok", f"vs. «{previa.name}»")
        self.t_med.poner(len(self.agendas), "info", "Ver agendas")
        fest = resumen["festivos"]
        self.t_fest.poner(len(fest), "info", ", ".join(analisis.fecha_corta(f) for f in fest) or "Ninguno")

        filas, colores, ayudas = [], {}, {}
        for r, srv in enumerate(COBERTURA):
            fila = [analisis.SERVICIOS.get(srv, srv)]
            for k, dia in enumerate(self.datos["dias"]):
                i = k + 1
                if srv == "DENSITO":
                    n = len(self.datos["densito"][i])
                else:
                    n = len({t[4] for t in self.datos["turnos"][i] if t[0] == srv})
                faltan = [fr for f, fr, s, _ in pend if s == srv and f == dia]
                fila.append(n or "")
                if faltan:
                    colores[(r, k + 1)] = "#FFF4D6"
                    ayudas[(r, k + 1)] = "Sin médico: " + ", ".join(faltan)
            filas.append(fila)
        llenar(self.t_cob, filas, lambda r, c: colores.get((r, c)))
        for (r, c), ayuda in ayudas.items():
            self.t_cob.item(r, c).setToolTip(ayuda)
        for r in range(self.t_cob.rowCount()):
            for c in range(1, self.t_cob.columnCount()):
                self.t_cob.item(r, c).setTextAlignment(Qt.AlignCenter)
        self.t_cob.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)

        n_alertas = len(alertas) + len(revision)
        for i, n in ((ALERTAS, n_alertas), (CAMBIOS, len(cambios or []))):
            self.nav.item(i).setText(f"{SECCIONES[i]}   ·  {n}" if n else SECCIONES[i])
        if cruces or resumen["desconocidos"]:
            partes = []
            if cruces:
                partes.append(f"{len(cruces)} turno(s) que se cruzan en el cuadro")
            if resumen["desconocidos"]:
                partes.append("códigos no reconocidos (no se publican): " + ", ".join(resumen["desconocidos"]))
            self.avisar("Revisa antes de publicar: " + "; ".join(partes) + ".", "warn")
        else:
            self.avisar("")

    def _mostrar_previa(self):
        self._previa_pendiente = False
        actual = self.previa.currentIndex()
        self.previa.clear()
        tema = analisis.paleta_tema(self.wb)
        for hoja, titulo in PESTANAS_PREVIA.items():
            if hoja in self.wb.sheetnames:
                self.previa.addTab(tabla_hoja(self.wb[hoja], 1 if hoja == "Cambios" else 3, tema), titulo)
        self.previa.setCurrentIndex(max(actual, 0))

    def _mostrar_alertas(self, alertas, pendientes, revision):
        llenar(self.t_revision, [(analisis.fecha_corta(f) if f else "Toda la semana", m, t)
                                 for f, m, t in revision] or [("", "", "Sin problemas en el cuadro.")])
        llenar(self.t_alertas, [(analisis.fecha_corta(f), m, t) for f, m, t in alertas]
               or [("", "", "Se cumplen las reglas de descanso.")])
        que = {"disp": "Disponibilidad sin asignar", "turno": "Turno sin médico", "remoto": "Remoto sin médico"}
        llenar(self.t_pendientes, [(analisis.fecha_corta(f), analisis.SERVICIOS.get(s, s), fr, que[t])
                                   for f, fr, s, t in pendientes] or [("", "", "", "Todas las franjas tienen médico.")])

    def _mostrar_equidad(self, equidad):
        claves = ["medico", "semana", "jornada", "mes", "meta", "pendiente", "noches", "domingos",
                  "festivos", "findes_libres", "disponibilidades"]

        def color(r, c):
            e = equidad[r]
            if c == 1 and e["jornada"] and e["semana"] > e["jornada"]:
                return "#FDE8E8"
            if c == 5 and e["pendiente"]:
                return "#FFF4D6" if e["pendiente"] > 0 else "#E8F1F8"
        llenar(self.t_equidad, [[e[k] for k in claves] for e in equidad], color, orden=0)

    def _mostrar_meses(self, *_):
        self.pila_eq.setCurrentIndex(self.modo_eq.currentIndex())
        if self.modo_eq.currentIndex() == 0 or self.ws is None:
            return
        otros = [r for r in (self.cfg.value("otros_meses", []) or []) if Path(r) != Path(self.ruta)]
        hojas, fallidos = [(self.ws.title.replace("Cambios ", ""), self.ws)], []
        with ocupado():
            for r in otros:
                try:
                    ws = gs.cargar_mensual(r)
                    hojas.append((ws.title.replace("Cambios ", ""), ws))
                except gs.ErrorPlanilla:
                    fallidos.append(Path(r).name)
        filas = analisis.equidad_meses(hojas)
        llenar(self.t_meses, [[f["medico"], len(f["meses"])] + [f[k] for k in analisis.METRICAS_MES] for f in filas],
               lambda r, c: ("#FFF4D6" if filas[r]["pendiente"] > 0 else "#E8F1F8")
               if c == 4 and filas[r]["pendiente"] else None, orden=0)
        texto = "Meses: " + ", ".join(e for e, _ in hojas)
        if len(hojas) == 1:
            texto += "  ·  Elige otros cuadros para sumar más meses."
        if fallidos:
            texto += "  ·  No se pudieron leer: " + ", ".join(fallidos)
        self.lbl_meses.setText(texto)

    def _mostrar_agendas(self):
        self.agendas = analisis.agenda_por_medico(self.datos)
        actual = self.l_medicos.currentItem().text() if self.l_medicos.currentItem() else None
        self.l_medicos.clear()
        self.l_medicos.addItems(list(self.agendas))
        self._filtrar_medicos(self.buscar.text())
        items = self.l_medicos.findItems(actual, Qt.MatchExactly) if actual else []
        self.l_medicos.setCurrentItem(items[0] if items else self.l_medicos.item(0))

    def _filtrar_medicos(self, texto):
        t = texto.strip().lower()
        for i in range(self.l_medicos.count()):
            it = self.l_medicos.item(i)
            it.setHidden(bool(t) and t not in it.text().lower())

    def mostrar_agenda(self, nombre):
        if nombre in self.agendas:
            self.txt_agenda.setPlainText(analisis.texto_agenda(nombre, self.agendas[nombre], self.lunes))

    def _mostrar_cambios(self, previa, cambios):
        if not previa:
            self.lbl_cambios.setText("Aún no se ha guardado ninguna versión de esta semana.")
            self.txt_cambios.setPlainText("")
        elif not cambios:
            self.lbl_cambios.setText(f"Sin cambios respecto a «{previa.name}».")
            self.txt_cambios.setPlainText("")
        else:
            self.lbl_cambios.setText(f"{len(cambios)} cambio(s) respecto a «{previa.name}».")
            self.txt_cambios.setPlainText("\n".join(cambios))


def error_inesperado(tipo, valor, tb):
    texto = "".join(traceback.format_exception(tipo, valor, tb))
    try:
        with open(LOG, "a", encoding="utf-8") as f:
            f.write(f"\n--- {datetime.now():%Y-%m-%d %H:%M}\n{texto}")
    except OSError:
        pass
    QMessageBox.critical(None, "Algo salió mal",
                         f"Ocurrió un error inesperado:\n\n{valor}\n\nSe guardó el detalle en {LOG.name}.")


def preparar_app(argv):
    app = QApplication(argv)
    app.setApplicationName("Planilla Semanal")
    app.setWindowIcon(icono_app())
    app.setStyle("Fusion")
    # Siempre modo claro: con el sistema en oscuro, Qt pondría texto claro sobre nuestros fondos blancos
    app.styleHints().setColorScheme(Qt.ColorScheme.Light)
    app.setPalette(app.style().standardPalette())
    if sys.platform == "win32":
        app.setFont(QFont("Segoe UI", 10))
    app.setStyleSheet(ESTILO)
    return app


def main():
    app = preparar_app(sys.argv)
    sys.excepthook = error_inesperado
    v = Ventana()
    v.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
