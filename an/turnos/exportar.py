"""Salidas en Excel: cuadro mensual, planilla de puestos y reporte de horas y recargos. Nunca sobrescribe."""
from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from an.nucleo.xlsx import poner_texto

from . import recargos
from . import reglas
from .dominio import DIAS, Puesto
from .puestos import alertas_reparto, etiqueta, leyenda, por_persona

ABREV_NOVEDAD = {"Vacaciones": "VAC", "Incapacidad": "INC", "Licencia": "LIC", "Libre solicitado": "LIB"}
ENCABEZADO = PatternFill("solid", fgColor="1F3864")
AMARILLO = PatternFill("solid", fgColor="FFFF00")
GRIS = PatternFill("solid", fgColor="D9D9D9")
FESTIVO = PatternFill("solid", fgColor="F8CBAD")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def ruta_libre(carpeta, base):
    """`base.xlsx`, luego `base v2.xlsx`, `base v3.xlsx`…"""
    carpeta, n = Path(carpeta), 1
    while (r := carpeta / (f"{base}.xlsx" if n == 1 else f"{base} v{n}.xlsx")).exists():
        n += 1
    return r


def _encabezado(ws, fila, textos, col=1):
    for i, t in enumerate(textos):
        c = ws.cell(fila, col + i, t)
        c.fill, c.font, c.alignment = ENCABEZADO, Font(bold=True, color="FFFFFF"), CENTRO


def _ancho(ws, primera, resto, desde=2):
    ws.column_dimensions["A"].width = primera
    for i in range(desde, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = resto


def _hoja_cuadro(ws, ctx, cuadro, fechas, comparar=None, puestos=None, detalle=None, contadores=None):
    """`detalle` ({colab_id: recargos.reporte}): horas, de más/menos y balance por turno como fórmulas sobre las celdas del
    cuadro (hoja «Turnos» del mismo libro) y recargos como valores. `contadores` ([(puesto, días sin cubrir)]): una fila por
    puesto bajo el cuadro; la hoja lleva celdas «M·2»."""
    cols = [p for p in ctx.colabs if p.en_solver]
    n, ult = len(cols) + 1, get_column_letter(len(fechas) + 1)
    solver = [t for t in ctx.turnos.values() if t.activo and t.solver] if detalle is not None else []
    extra = ["De más", "De menos"] + [t.codigo for t in solver] + ["Horas equiv. recargos", "Dom/fest trabajados"] if detalle is not None else []
    contr = len(fechas) + 3
    _encabezado(ws, 1, ["Colaborador"] + [f"{DIAS[f.weekday()]}\n{f.day}" for f in fechas] + ["Horas"] + (["Contratadas"] + extra if detalle is not None else []))
    for i, f in enumerate(fechas, 2):
        if ctx.tipo_dia(f) == "DF":
            ws.cell(1, i).fill = PatternFill("solid", fgColor="C55A11")
    for r, p in enumerate(cols, 2):
        poner_texto(ws.cell(r, 1), p.nombre).font = Font(bold=True)
        horas = 0
        for i, f in enumerate(fechas, 2):
            cod = cuadro.get((p.id, f), "L")
            nov = ctx.novedad(p.id, f)
            texto = ABREV_NOVEDAD[nov.tipo] if nov and cod == "L" else cod
            if puestos is not None and cod in ctx.turnos and (p.id, f) in puestos:
                texto = etiqueta(cod, puestos[p.id, f])
            c = poner_texto(ws.cell(r, i), texto)
            c.alignment = CENTRO
            if cod in ctx.turnos:
                c.fill = PatternFill("solid", fgColor=ctx.turnos[cod].color)
                horas += ctx.turnos[cod].horas
            elif nov:
                c.fill = GRIS
            if comparar is not None and comparar.get((p.id, f), "L") != cod:
                c.fill = AMARILLO
        if detalle is None:
            ws.cell(r, len(fechas) + 2, horas)
            continue
        fila, rep = f"$B{r}:${ult}{r}", detalle[p.id]
        ws.cell(r, len(fechas) + 2, f"=SUMPRODUCT(COUNTIF({fila},Turnos!$A$2:$A${len(ctx.turnos) + 1}),Turnos!$C$2:$C${len(ctx.turnos) + 1})")
        h, c = f"{get_column_letter(len(fechas) + 2)}{r}", f"{get_column_letter(contr)}{r}"
        ws.cell(r, contr, round(rep["contratadas"], 1))
        ws.cell(r, contr + 1, f"=MAX(0,{h}-{c})")
        ws.cell(r, contr + 2, f"=MAX(0,{c}-{h})")
        for k, t in enumerate(solver):
            ws.cell(r, contr + 3 + k, f'=COUNTIF({fila},"{t.codigo}")')
        ws.cell(r, contr + 3 + len(solver), round(rep["equivalentes"], 1))
        ws.cell(r, contr + 4 + len(solver), rep["festivos_trabajados"])
    base = len(cols) + 3
    faltantes = reglas.faltantes_periodo(ctx, cuadro, fechas)
    for k, (fr, nombre) in enumerate((("M", "Mañana"), ("T", "Tarde"), ("N", "Noche"))):
        ws.cell(base + k, 1, f"Cobertura {nombre.lower()}").font = Font(bold=True)
        for i, f in enumerate(fechas, 2):
            peor, rangos = faltantes[f][fr]
            c = ws.cell(base + k, i, "ok" if not peor else f"−{peor}")
            c.alignment = CENTRO
            if peor:
                c.fill = PatternFill("solid", fgColor="FF7C80")
    fila = base + 3
    activos = [t for t in ctx.turnos.values() if t.activo and t.solver]
    if activos:  # contador por turno: personas con ese código cada día (en la hoja de puestos «M·2» cuenta como M); rojo si la franja que cubre queda corta
        ws.cell(fila, 1, "Contador por turno").font = Font(bold=True)
        for k, t in enumerate(activos, 1):
            ws.cell(fila + k, 1, f"{t.codigo} {t.nombre}")
            for i, f in enumerate(fechas, 2):
                col = get_column_letter(i)
                rango = f"{col}$2:{col}${n}"
                c = ws.cell(fila + k, i, f'=COUNTIF({rango},"{t.codigo}")+COUNTIF({rango},"{t.codigo}·*")')
                c.alignment = CENTRO
                if any(faltantes[f][fr][0] for fr in t.cubre):
                    c.fill = PatternFill("solid", fgColor="FF7C80")
        fila += len(activos) + 2
    if contadores:  # contador por puesto: personas-día en cada puesto; rojo si la app lo reporta sin cubrir
        ws.cell(fila, 1, "Contador por puesto").font = Font(bold=True)
        for k, (q, sin) in enumerate(contadores, 1):
            ws.cell(fila + k, 1, f"{q.numero} {q.nombre}" if q.numero else "A Apoyo")
            for i, f in enumerate(fechas, 2):
                col = get_column_letter(i)
                c = ws.cell(fila + k, i, f'=COUNTIF({col}$2:{col}${n},"*·{q.numero or "A"}")')
                c.alignment = CENTRO
                if f in sin:
                    c.fill = PatternFill("solid", fgColor="FF7C80")
    _ancho(ws, 22, 5)
    ws.column_dimensions[get_column_letter(len(fechas) + 2)].width = 8
    for i in range(contr, contr + len(extra) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11 if detalle is not None else 5
    ws.freeze_panes = "B2"


def exportar_cuadro(carpeta, mes_txt, ctx, cuadro, original, fechas, diferencias, filas_puestos=(), catalogo=None, detalle=None):
    """Hojas: Cuadro (de `cambios`, celdas modificadas en amarillo; con `detalle`, columnas de horas y contadores por turno),
    Cuadro con puestos (con contador por puesto), Original, Cambios (diferencias) y Turnos (catálogo que leen las fórmulas)."""
    wb = Workbook()
    _hoja_cuadro(wb.active, ctx, cuadro, fechas, comparar=original, detalle=detalle)
    wb.active.title = "Cuadro"
    if filas_puestos:
        con = wb.create_sheet("Cuadro con puestos")
        cat = catalogo or ctx.puestos
        activos = sorted((q for q in cat.values() if q.activo), key=lambda q: q.numero)
        falta = {q.id: {a.fecha for a in alertas_reparto(ctx, filas_puestos, fechas, cat) if f": {q.nombre} sin cubrir" in a.texto} for q in activos}
        _hoja_cuadro(con, ctx, cuadro, fechas, puestos=por_persona(filas_puestos, {i: p.numero for i, p in cat.items()}),
                     contadores=[(q, falta[q.id]) for q in activos] + [(Puesto(0, 0, "Apoyo"), set())] if detalle is not None else None)
        poner_texto(con.cell(con.max_row + 2, 1), leyenda(cat))
    _hoja_cuadro(wb.create_sheet("Original"), ctx, original, fechas)
    ws = wb.create_sheet("Cambios")
    _encabezado(ws, 1, ["Colaborador", "Fecha", "Antes", "Después"])
    nombres = {p.id: p.nombre for p in ctx.colabs}
    for r, (cid, f, antes, ahora) in enumerate(diferencias, 2):
        for c, v in enumerate((nombres.get(cid, cid), f.strftime("%d/%m/%Y"), antes, ahora), 1):
            poner_texto(ws.cell(r, c), v)
    _ancho(ws, 26, 14)
    if detalle is not None:  # lo leen las fórmulas de horas del cuadro
        tur = wb.create_sheet("Turnos")
        _encabezado(tur, 1, ["Código", "Nombre", "Horas"])
        for r, t in enumerate(ctx.turnos.values(), 2):
            for c, v in enumerate((t.codigo, t.nombre, round(t.horas, 2)), 1):
                poner_texto(tur.cell(r, c), v) if c < 3 else tur.cell(r, c, v)
        _ancho(tur, 10, 14)
    ruta = ruta_libre(carpeta, f"CUADRO {mes_txt.upper()}")
    wb.save(ruta)
    return ruta


def _hm(minutos):
    return f"{minutos // 60 % 24:02d}:{minutos % 60:02d}"


def exportar_puestos(carpeta, mes_txt, ctx, cuadro, fechas, filas, capacitacion, catalogo=None):
    """Una hoja por semana: por día, quién ocupa cada puesto (con su horario y grupo de pausa), el apoyo y la capacitación."""
    wb = Workbook()
    wb.remove(wb.active)
    nombres = {p.id: p.nombre for p in ctx.colabs}
    cat = catalogo or ctx.puestos
    usados = sorted({f.puesto_id for f in filas if f.puesto_id and f.puesto_id in cat} | {i for i, p in cat.items() if p.activo},
                    key=lambda i: cat[i].numero)
    cap = {}
    for f, cid in capacitacion:
        cap.setdefault(f, []).append(nombres[cid])
    celdas = {}
    for a in filas:
        txt = f"{nombres.get(a.colab_id, a.colab_id)} ({_hm(a.desde)}–{_hm(a.hasta % 1440 or 1440)}) (P{a.pausa})"
        celdas.setdefault((a.fecha, a.puesto_id), []).append(txt)
    for lunes in sorted({f - timedelta(days=f.weekday()) for f in fechas}):
        ws = wb.create_sheet(f"Semana {lunes:%d-%m}")
        _encabezado(ws, 1, ["Día"] + [f"{cat[i].numero} · {cat[i].nombre}" for i in usados] + ["Apoyo", "Capacitación"])
        for r, f in enumerate((lunes + timedelta(days=i) for i in range(7)), 2):
            c = ws.cell(r, 1, f"{DIAS[f.weekday()]} {f:%d/%m}")
            c.font = Font(bold=True)
            if ctx.tipo_dia(f) == "DF":
                c.fill = FESTIVO
            if f not in fechas:
                continue
            fila = ["\n".join(celdas.get((f, i), [])) for i in usados] + ["\n".join(celdas.get((f, 0), []))] + [", ".join(cap.get(f, []))]
            for i, v in enumerate(fila, 2):
                poner_texto(ws.cell(r, i), v).alignment = CENTRO
        ws.cell(10, 1, "(P1)/(P2): grupo de pausa — el grupo 1 sale primero y el 2 cubre; se alternan cada día.")
        _ancho(ws, 14, 26)
        ws.freeze_panes = "B2"
    ruta = ruta_libre(carpeta, f"PUESTOS {mes_txt.upper()}")
    wb.save(ruta)
    return ruta


CATEGORIAS = [(t, n, f) for t in ("ordinaria", "extra") for f in ("hábil", "festivo") for n in ("diurna", "nocturna")]


def exportar_horas(carpeta, mes_txt, ctx, tarifas, cuadro, fechas, ventana_sem):
    wb = Workbook()
    ws = wb.active
    ws.title = "Horas y recargos"
    nombre_cat = lambda c: f"{'Extra' if c[0] == 'extra' else 'Ordinaria'} {c[1]} {'dom/fest' if c[2] == 'festivo' else 'hábil'}"
    _encabezado(ws, 1, ["Colaborador", "Contratadas", "Trabajadas", "Pendientes", "Adicionales"]
                + [nombre_cat(c) for c in CATEGORIAS] + ["Dom/fest trabajados", "Habitual", "Horas equivalentes de recargo"])
    sem = wb.create_sheet("Semanas")
    lunes = sorted({f - timedelta(days=f.weekday()) for f in fechas})
    _encabezado(sem, 1, ["Colaborador"] + [f"Sem {l:%d/%m}" for l in lunes])
    for r, p in enumerate((p for p in ctx.colabs if p.en_solver), 2):
        rep = recargos.reporte(ctx, tarifas, p, cuadro, fechas, ventana_sem)
        fila = [p.nombre] + [round(rep[k], 1) for k in ("contratadas", "trabajadas", "pendientes", "adicionales")] \
            + [round(rep["cat"].get(c, 0), 1) for c in CATEGORIAS] \
            + [rep["festivos_trabajados"], "Sí" if rep["habitual"] else "No", round(rep["equivalentes"], 1)]
        for c, v in enumerate(fila, 1):
            poner_texto(ws.cell(r, c), v)
        poner_texto(sem.cell(r, 1), p.nombre)
        for c, l in enumerate(lunes, 2):
            sem.cell(r, c, round(rep["semanas"].get(l, 0), 1))
    _ancho(ws, 24, 14)
    _ancho(sem, 24, 12)
    ws.freeze_panes = "B2"
    nota = wb.create_sheet("Régimen")
    f0 = fechas[0]
    filas = [("Régimen de horas", f"Art. 165 CST: promedio en ventanas de {ventana_sem} semanas contra la jornada contratada."),
             ("Ventana", f"Contada desde el primer día del periodo; no cruza periodos."),
             ("Pendientes / adicionales", "Déficit o exceso de la ventana frente a lo contratado (descontando novedades)."),
             ("Horas extra", "Se toman de las últimas horas trabajadas de la ventana."),
             ("Dominical/festivo", f"{recargos.tarifa(tarifas, 'dominical', f0)} %"),
             ("Nocturno", f"{recargos.tarifa(tarifas, 'nocturno', f0)} % desde las {recargos.tarifa(tarifas, 'noche_inicio', f0) // 60}:00"),
             ("Extra diurna / nocturna", f"{recargos.tarifa(tarifas, 'extra_diurna', f0)} % / {recargos.tarifa(tarifas, 'extra_nocturna', f0)} %"),
             ("Horas equivalentes", "Ordinarias: horas × recargo. Extras: horas × (1 + recargo). Los recargos se suman."),
             ("Aviso", "Cifras de apoyo para nómina; confirmar con RR. HH./jurídico.")]
    for r, (a, b) in enumerate(filas, 1):
        nota.cell(r, 1, a).font = Font(bold=True)
        nota.cell(r, 2, b)
    _ancho(nota, 28, 90)
    ruta = ruta_libre(carpeta, f"HORAS Y RECARGOS {mes_txt.upper()}")
    wb.save(ruta)
    return ruta
